"""Run INSIDE a disposable pinned sandbox container, never against production.

Builds tiny synthetic fixtures and verifies a pinned Apache-licensed upstream MSG
attachment fixture. No user document is uploaded by this harness.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import signal
import subprocess
import tempfile
import time
import urllib.error
import urllib.request
import zipfile
from email.message import EmailMessage
from pathlib import Path


MARKER = "KnowledgeFS Golden Evidence"


def pdf_fixture() -> bytes:
    stream = f"BT /F1 12 Tf 20 100 Td ({MARKER}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    body = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(body))
        body.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    startxref = len(body)
    body.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        body.extend(f"{offset:010} 00000 n \n".encode())
    body.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{startxref}\n%%EOF\n".encode()
    )
    return bytes(body)


def multipart(filename: str, body: bytes) -> tuple[bytes, str]:
    boundary = "kfs-golden-" + os.urandom(12).hex()
    fields = bytearray()
    for name, value in (
        ("strategy", "fast"),
        ("coordinates", "true"),
        ("include_slide_notes", "true"),
    ):
        fields.extend(
            f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n{value}\r\n'.encode()
        )
    fields.extend(
        f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="{filename}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode()
    )
    fields.extend(body)
    fields.extend(f"\r\n--{boundary}--\r\n".encode())
    return bytes(fields), f"multipart/form-data; boundary={boundary}"


def has_evidence(value, marker: str = MARKER) -> bool:
    return isinstance(value, list) and marker in " ".join(
        str(item.get("text", "")) for item in value if isinstance(item, dict)
    )


def msg_fixture(path: Path | None) -> tuple[bytes, str]:
    target = path or Path(__file__).with_name("fixtures") / "fake-email-attachment.msg"
    if target.stat().st_size > 1024 * 1024:
        raise ValueError("MSG golden fixture must be at most 1 MiB")
    body = target.read_bytes()
    if path is None:
        if (
            hashlib.sha256(body).hexdigest()
            != "92f65236e7eae301ea6f70a85f38cd5a9fae9f807d17fc35c5b5dbcf0f82a8ec"
        ):
            raise ValueError("Upstream MSG fixture digest mismatch")
        return body, "Hey this is a fake attachment!"
    return body, MARKER


def convert(command: list[str]) -> None:
    # Fixture construction runs outside a request. Use the image-captured original
    # tools only for these tiny trusted inputs; HTTP parsing still uses the wrappers.
    manifest = Path("/opt/kfs-sandbox/converters.json")
    if manifest.is_file():
        command = [json.loads(manifest.read_text())[command[0]], *command[1:]]
    process = subprocess.Popen(
        command,
        start_new_session=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        code = process.wait(timeout=45)
        if code != 0:
            raise RuntimeError("fixture_conversion_failed")
    finally:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        process.wait()


def create_fixtures(directory: Path) -> dict[str, bytes]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    document = Document()
    document.add_paragraph(MARKER)
    document.save(directory / "sample.docx")
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = MARKER
    slides.save(directory / "sample.pptx")
    workbook = Workbook()
    workbook.active.append(["evidence"])
    workbook.active.append([MARKER])
    workbook.save(directory / "sample.xlsx")
    profile = (directory / "office-profile").as_uri()
    for source, target in (
        ("docx", "doc"),
        ("pptx", "ppt"),
        ("xlsx", "xls"),
        ("docx", "odt"),
    ):
        convert(
            [
                "soffice",
                f"-env:UserInstallation={profile}",
                "--headless",
                "--convert-to",
                target,
                "--outdir",
                str(directory),
                str(directory / f"sample.{source}"),
            ]
        )
    html = directory / "sample.html"
    html.write_text(f"<html><body><p>{MARKER}</p></body></html>")
    convert(["pandoc", str(html), "-o", str(directory / "sample.epub")])
    fixtures = {
        f"sample.{extension}": (directory / f"sample.{extension}").read_bytes()
        for extension in ("doc", "docx", "ppt", "pptx", "xls", "xlsx", "odt", "epub")
    }
    fixtures["sample.pdf"] = pdf_fixture()
    fixtures["sample.rtf"] = (r"{\rtf1\ansi " + MARKER + "}").encode()
    message = EmailMessage()
    message["Subject"] = "Golden attachment test"
    message.set_content("The evidence is in the attachment.")
    message.add_attachment(
        fixtures["sample.docx"],
        maintype="application",
        subtype="octet-stream",
        filename="attached.docx",
    )
    fixtures["sample.eml"] = message.as_bytes()
    return fixtures


def post(url: str, name: str, payload: bytes):
    body, content_type = multipart(name, payload)
    request = urllib.request.Request(
        url, body, {"Content-Type": content_type, "Accept": "application/json"}
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            raw = response.read(32 * 1024 * 1024 + 1)
            if len(raw) > 32 * 1024 * 1024:
                raise RuntimeError("response_bytes")
            return response.status, json.loads(raw), len(raw)
    except urllib.error.HTTPError as error:
        return error.code, None, 0


def run_gate(url: str, custom_msg_fixture: Path | None) -> dict:
    # Refuse accidental production runs, even if a user copies a production URL.
    if url not in {
        "http://127.0.0.1:8000/general/v0/general",
        "http://localhost:8000/general/v0/general",
    }:
        raise ValueError(
            "Golden requests must target the disposable container's loopback address"
        )
    report = {"passed": [], "failures": [], "missing_fixtures": []}
    with tempfile.TemporaryDirectory(prefix="kfs-golden-") as directory:
        fixtures = create_fixtures(Path(directory))
        msg_body, msg_marker = msg_fixture(custom_msg_fixture)
        fixtures["sample.msg"] = msg_body
        for name, body in fixtures.items():
            started = time.monotonic()
            status, result, response_bytes = post(url, name, body)
            if status == 200 and has_evidence(
                result, msg_marker if name == "sample.msg" else MARKER
            ):
                report["passed"].append(
                    {
                        "format": name.rsplit(".", 1)[-1],
                        "milliseconds": round((time.monotonic() - started) * 1000),
                        "input_bytes": len(body),
                        "response_bytes": response_bytes,
                    }
                )
            else:
                report["failures"].append(
                    {
                        "format": name.rsplit(".", 1)[-1],
                        "status": status,
                        "reason": "missing_evidence_or_failed",
                    }
                )
        # Tiny sparse extent fixture; admission must reject it BEFORE pandas allocation.
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr(
                "xl/worksheets/sheet1.xml", '<worksheet><c r="XFD1048576"/></worksheet>'
            )
        nested = EmailMessage()
        nested["Subject"] = "Bounded rejection fixture"
        nested.set_content("Body")
        nested.add_attachment(
            stream.getvalue(),
            maintype="application",
            subtype="octet-stream",
            filename="unsafe.xlsx",
        )
        status, _, _ = post(url, "nested.eml", nested.as_bytes())
        if status != 422:
            report["failures"].append(
                {
                    "format": "eml-nested-xlsx",
                    "status": status,
                    "reason": "admission_did_not_reject",
                }
            )
    report["gate"] = (
        "passed"
        if not report["failures"] and not report["missing_fixtures"]
        else "blocked"
    )
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--msg-fixture", type=Path)
    args = parser.parse_args()
    result = run_gate("http://127.0.0.1:8000/general/v0/general", args.msg_fixture)
    print(json.dumps(result, indent=2))
    raise SystemExit(0 if result["gate"] == "passed" else 1)
