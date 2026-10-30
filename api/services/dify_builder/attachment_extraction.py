"""Opt-in, local-only Builder context extraction; no Knowledge or image pipeline.

The caller supplies the native storage stream after authorizing its UploadFile.
No parser follows links, performs OCR, writes derivatives or invokes providers.
"""

import hashlib
import io
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from typing import Literal

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_FILE_CHARS = 20_000
MAX_ZIP_BYTES = 20 * 1024 * 1024


@dataclass(frozen=True)
class ExtractedContext:
    status: Literal["ready", "truncated", "rejected"]
    text: str = ""
    digest: str = ""
    reason: str = ""


class _RejectedError(ValueError):
    pass


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader, apply_configuration
    from pypdf.errors import LimitReachedError
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject, StreamObject

    stream_limit = 2 * 1024 * 1024
    safe_filters = {"/FlateDecode", "/ASCIIHexDecode", "/ASCII85Decode", "/LZWDecode", "/RunLengthDecode"}
    decoded_bytes = 0
    visited: set[int] = set()

    def inspect(value, depth=0):
        nonlocal decoded_bytes
        if depth > 32 or len(visited) > 512:
            raise _RejectedError("pdf_content_limit")
        if isinstance(value, IndirectObject):
            value = value.get_object()
        if id(value) in visited:
            return
        visited.add(id(value))
        if isinstance(value, StreamObject):
            if value.get("/Subtype") == "/Image":
                return  # Image streams are never decoded by this text profile.
            if value.get("/Subtype") == "/Form":
                raise _RejectedError("unsupported_pdf_content")
            filters = value.get("/Filter", [])
            if not isinstance(filters, (list, tuple)):
                filters = [filters]
            if any(str(item) not in safe_filters for item in filters):
                raise _RejectedError("unsupported_pdf_filter")
            size = len(value.get_data())
            decoded_bytes += size
            if size > stream_limit or decoded_bytes > MAX_ZIP_BYTES:
                raise _RejectedError("pdf_content_limit")
        if isinstance(value, DictionaryObject):
            for key, child in value.items():
                if key not in {"/Parent", "/Metadata"}:
                    inspect(child, depth + 1)
        elif isinstance(value, ArrayObject):
            for child in value:
                inspect(child, depth + 1)

    # Public context-local decoder settings cover content, fonts/CMaps and
    # compressed object metadata without changing process-global parser policy.
    with apply_configuration(
        maximum_declared_stream_length=stream_limit,
        array_based_stream_maximum_output_length=stream_limit,
        zlib_maximum_output_length=stream_limit,
        lzw_maximum_output_length=stream_limit,
        run_length_maximum_output_length=stream_limit,
        brotli_maximum_output_length=stream_limit,
        stream_decoding_work_maximum_length=MAX_ZIP_BYTES,
        stream_parsing_work_maximum_count=100_000,
        array_maximum_element_count=2000,
        dictionary_maximum_element_count=2000,
        page_tree_maximum_entries=200,
        xform_maximum_invocations_per_extraction=0,
        jbig2dec_binary=None,
    ):
        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise _RejectedError("encrypted_file")
            if len(reader.pages) > 100:
                raise _RejectedError("page_limit")
            parts: list[str] = []
            length = 0
            for page in reader.pages:
                resources = page.get("/Resources", {})
                if isinstance(resources, DictionaryObject):
                    xobjects = resources.get("/XObject", {})
                    for reference in xobjects.values():
                        if reference.get_object().get("/Subtype") != "/Image":
                            raise _RejectedError("unsupported_pdf_content")
                inspect(resources)
                inspect(page.get("/Contents"))
                text = page.extract_text() or ""
                remaining = MAX_FILE_CHARS + 1 - length
                parts.append(text[:remaining])
                length += len(parts[-1])
                if length > MAX_FILE_CHARS:
                    break
            return "\n".join(parts)
        except LimitReachedError as exc:
            raise _RejectedError("pdf_content_limit") from exc


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        raise _RejectedError("encrypted_file")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(entry.file_size for entry in entries) > MAX_ZIP_BYTES:
            raise _RejectedError("zip_expansion_limit")
        if any(entry.flag_bits & 1 for entry in entries):
            raise _RejectedError("encrypted_file")
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        if len(book.worksheets) > 10:
            raise _RejectedError("sheet_limit")
        parts: list[str] = []
        cells = 0
        length = 0
        for sheet in book.worksheets:
            if (sheet.max_row or 0) * (sheet.max_column or 0) + cells > 10_000:
                raise _RejectedError("cell_limit")
            parts.append(f"[{sheet.title}]\n")
            length += len(parts[-1])
            for row in sheet.iter_rows(values_only=True):
                cells += len(row)
                if cells > 10_000:
                    raise _RejectedError("cell_limit")
                line = "\t".join("" if cell is None else str(cell) for cell in row) + "\n"
                remaining = MAX_FILE_CHARS + 1 - length
                parts.append(line[:remaining])
                length += len(parts[-1])
                if length > MAX_FILE_CHARS:
                    return "".join(parts)
        return "".join(parts)
    finally:
        book.close()


def extract_local_context(
    extension: str, chunks: Iterable[bytes], *, max_bytes: int = MAX_FILE_BYTES
) -> ExtractedContext:
    """Read at most 10 MiB; extract at most 20k characters or reject honestly."""
    extension = extension.lower().lstrip(".")
    if extension not in {"txt", "md", "markdown", "csv", "pdf", "xlsx"}:
        return ExtractedContext("rejected", reason="unsupported_format")
    data = bytearray()
    try:
        for chunk in chunks:
            if len(data) + len(chunk) > min(MAX_FILE_BYTES, max_bytes):
                return ExtractedContext("rejected", reason="byte_limit")
            data.extend(chunk)
        digest = hashlib.sha256(data).hexdigest()
        raw = bytes(data)
        if extension == "pdf":
            text = _pdf(raw)
        elif extension == "xlsx":
            text = _xlsx(raw)
        else:
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                return ExtractedContext("rejected", digest=digest, reason="unsupported_encoding")
        if not text.strip():
            return ExtractedContext("rejected", digest=digest, reason="no_extractable_text")
        truncated = len(text) > MAX_FILE_CHARS
        return ExtractedContext(
            "truncated" if truncated else "ready", text[:MAX_FILE_CHARS], digest, "character_limit" if truncated else ""
        )
    except _RejectedError as exc:
        return ExtractedContext("rejected", reason=str(exc))
    except Exception:
        return ExtractedContext("rejected", reason="malformed_file")
    finally:
        close = getattr(chunks, "close", None)
        if close is not None:
            close()
