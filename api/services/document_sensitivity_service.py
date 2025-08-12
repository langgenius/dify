import logging
import os
import tempfile
from typing import Any, Optional

import openpyxl
import pypdf
from docx import Document as DocxDocument
from pptx import Presentation

from services.errors.file import SensitiveDocumentError

logger = logging.getLogger(__name__)


class DocumentSensitivityService:
    """Service for checking document sensitivity levels from metadata."""

    @classmethod
    def check_document_sensitivity(
        cls, extension: str, content: bytes, blocked_levels: Optional[list[str]] = None
    ) -> Optional[str]:
        if blocked_levels is None:
            blocked_levels = []
        blocked_levels_lower = {level.lower() for level in blocked_levels}

        try:
            with tempfile.NamedTemporaryFile(suffix=extension, delete=False) as temp_file:
                temp_file.write(content)
                temp_file.flush()
                temp_path = temp_file.name

            try:
                metadata = cls._extract_metadata(temp_path, extension)

                detected_sensitivity = cls._check_metadata_sensitivity(metadata, blocked_levels_lower)

                if detected_sensitivity:
                    raise SensitiveDocumentError(f"sensitive info detected : {detected_sensitivity}")

                return detected_sensitivity

            finally:
                # Clean up temporary file
                if os.path.exists(temp_path):
                    os.unlink(temp_path)

        except SensitiveDocumentError:
            raise SensitiveDocumentError()
        except Exception as e:
            # Log the error but don't block upload for metadata extraction failure
            logger.warning("Failed to extract metadata from file %s", e)
            return None

    @classmethod
    def _extract_metadata(cls, file_path: str, extension: str) -> dict[str, Any]:
        """Extract metadata from document based on file type."""
        metadata: dict[str, Any] = {}
        try:
            if extension.lower() in {".docx", ".doc"}:
                metadata.update(cls._extract_word_metadata(file_path))
            elif extension.lower() == ".pdf":
                metadata.update(cls._extract_pdf_metadata(file_path))
            elif extension.lower() in {".xlsx", ".xls"}:
                metadata.update(cls._extract_excel_metadata(file_path))
            elif extension.lower() == ".pptx":
                metadata.update(cls._extract_powerpoint_metadata(file_path))
            elif extension.lower() in {".txt", ".md", ".markdown", ".mdx", ".csv"}:
                metadata.update(cls._extract_text_metadata(file_path, extension))
        except Exception as e:
            logger.warning("Failed to extract %s metadata: %s", extension, e)

        return metadata

    @classmethod
    def _extract_word_metadata(cls, file_path: str) -> dict[str, Any]:
        """Extract metadata from Word documents."""
        metadata = {}
        try:
            doc = DocxDocument(file_path)
            # Extract core properties
            core_props = doc.core_properties
            if hasattr(core_props, "category") and core_props.category:
                metadata["category"] = core_props.category.lower()
            if hasattr(core_props, "subject") and core_props.subject:
                metadata["subject"] = core_props.subject.lower()
            if hasattr(core_props, "comments") and core_props.comments:
                metadata["comments"] = core_props.comments.lower()
            if hasattr(core_props, "keywords") and core_props.keywords:
                metadata["keywords"] = core_props.keywords.lower()
            # Extract custom properties (including MIP labels)
            try:
                if hasattr(doc, "custom_doc_props"):
                    for prop in doc.custom_doc_props.props:
                        prop_name = prop.name.lower()
                        prop_value = str(prop.value).lower() if prop.value else ""
                        # Check for Microsoft Information Protection labels
                        if "msip_label" in prop_name and "_name" in prop_name:
                            metadata["mip_label"] = prop_value
                        elif "sensitivity" in prop_name:
                            metadata["sensitivity"] = prop_value
                        elif "classification" in prop_name:
                            metadata["classification"] = prop_value
                        metadata[prop_name] = prop_value
            except Exception:
                pass
        except Exception:
            pass
        return metadata

    @classmethod
    def _extract_pdf_metadata(cls, file_path: str) -> dict[str, Any]:
        """Extract metadata from PDF documents."""
        metadata = {}
        try:
            with open(file_path, "rb") as file:
                pdf_reader = pypdf.PdfReader(file)
                if pdf_reader.metadata:
                    for key, value in pdf_reader.metadata.items():
                        if value:
                            metadata[key.lower()] = str(value).lower()
        except Exception:
            pass
        return metadata

    @classmethod
    def _extract_excel_metadata(cls, file_path: str) -> dict[str, Any]:
        """Extract metadata from Excel documents."""
        metadata = {}
        try:
            workbook = openpyxl.load_workbook(file_path)
            # Extract core properties
            if hasattr(workbook, "properties"):
                props = workbook.properties
                if hasattr(props, "category") and props.category:
                    metadata["category"] = props.category.lower()
                if hasattr(props, "subject") and props.subject:
                    metadata["subject"] = props.subject.lower()
                if hasattr(props, "comments") and props.comments:
                    metadata["comments"] = props.comments.lower()
                if hasattr(props, "keywords") and props.keywords:
                    metadata["keywords"] = props.keywords.lower()
            # Extract custom properties
            if hasattr(workbook, "custom_doc_props"):
                for prop in workbook.custom_doc_props.props:
                    prop_name = prop.name.lower()
                    prop_value = str(prop.value).lower() if prop.value else ""
                    if "msip_label" in prop_name and "_name" in prop_name:
                        metadata["mip_label"] = prop_value
                    elif "sensitivity" in prop_name:
                        metadata["sensitivity"] = prop_value
                    elif "classification" in prop_name:
                        metadata["classification"] = prop_value
                    metadata[prop_name] = prop_value
        except Exception:
            pass
        return metadata

    @classmethod
    def _extract_powerpoint_metadata(cls, file_path: str) -> dict[str, Any]:
        """Extract metadata from PowerPoint documents."""
        metadata = {}
        try:
            pres = Presentation(file_path)
            # Extract core properties
            if hasattr(pres, "core_properties"):
                core_props = pres.core_properties
                if hasattr(core_props, "category") and core_props.category:
                    metadata["category"] = core_props.category.lower()
                if hasattr(core_props, "subject") and core_props.subject:
                    metadata["subject"] = core_props.subject.lower()
                if hasattr(core_props, "comments") and core_props.comments:
                    metadata["comments"] = core_props.comments.lower()
                if hasattr(core_props, "keywords") and core_props.keywords:
                    metadata["keywords"] = core_props.keywords.lower()
        except Exception:
            pass
        return metadata

    @classmethod
    def _extract_text_metadata(cls, file_path: str, extension: str) -> dict[str, Any]:
        """Extract metadata from text-based documents."""
        metadata: dict[str, Any] = {}
        try:
            # Read file content with encoding detection
            content = ""
            encodings_to_try = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
            for encoding in encodings_to_try:
                try:
                    with open(file_path, encoding=encoding) as file:
                        content = file.read()
                    break
                except UnicodeDecodeError:
                    continue
            if not content:
                return metadata
            # Extract metadata based on file type
            if extension.lower() in {".md", ".markdown", ".mdx"}:
                metadata.update(cls._extract_markdown_frontmatter(content))
            elif extension.lower() == ".csv":
                metadata.update(cls._extract_csv_headers(content))
            elif extension.lower() == ".txt":
                metadata.update(cls._extract_txt_metadata(content))
        except Exception as e:
            logger.warning("Failed to extract text metadata from %s: %s", extension, e)
        return metadata

    @classmethod
    def _extract_markdown_frontmatter(cls, content: str) -> dict[str, Any]:
        """Extract metadata from Markdown frontmatter (YAML/TOML)."""
        metadata = {}
        try:
            # Check for YAML frontmatter
            if content.startswith("---"):
                end_idx = content.find("\n---", 4)
                if end_idx != -1:
                    frontmatter = content[4:end_idx]
                    # Parse YAML-like frontmatter manually
                    for line in frontmatter.split("\n"):
                        line = line.strip()
                        if ":" in line and not line.startswith("#"):
                            key, value = line.split(":", 1)
                            key = key.strip().lower()
                            value = value.strip().strip("\"'").lower()
                            if key in [
                                "category",
                                "classification",
                                "sensitivity",
                                "tags",
                                "keywords",
                                "title",
                                "description",
                            ]:
                                metadata[key] = value
            # Check for TOML frontmatter
            elif content.startswith("+++"):
                end_idx = content.find("\n+++", 4)
                if end_idx != -1:
                    frontmatter = content[4:end_idx]
                    # Parse TOML-like frontmatter manually
                    for line in frontmatter.split("\n"):
                        line = line.strip()
                        if "=" in line and not line.startswith("#"):
                            key, value = line.split("=", 1)
                            key = key.strip().lower()
                            value = value.strip().strip("\"'").lower()
                            if key in [
                                "category",
                                "classification",
                                "sensitivity",
                                "tags",
                                "keywords",
                                "title",
                                "description",
                            ]:
                                metadata[key] = value
        except Exception:
            pass
        return metadata

    @classmethod
    def _extract_csv_headers(cls, content: str) -> dict[str, Any]:
        """Extract metadata from CSV headers."""
        metadata = {}
        try:
            lines = content.split("\n")
            if lines:
                # Get first line as headers
                headers = lines[0].lower()
                metadata["csv_headers"] = headers
                # Check for sensitive column names
                sensitive_columns = [
                    "password",
                    "ssn",
                    "social_security",
                    "credit_card",
                    "confidential",
                    "sensitive",
                    "classified",
                    "restricted",
                ]
                for sensitive in sensitive_columns:
                    if sensitive in headers:
                        metadata["sensitive_columns"] = sensitive
                        break
        except Exception:
            pass
        return metadata

    @classmethod
    def _extract_txt_metadata(cls, content: str) -> dict[str, Any]:
        """Extract metadata from plain text files."""
        metadata = {}
        try:
            lines = content.split("\n")
            # Look for metadata-like patterns in first few lines
            for i, line in enumerate(lines[:10]):  # Check first 10 lines
                line = line.strip().lower()
                # Look for key: value patterns
                if ":" in line:
                    parts = line.split(":", 1)
                    if len(parts) == 2:
                        key = parts[0].strip()
                        value = parts[1].strip()
                        if key in [
                            "classification",
                            "sensitivity",
                            "category",
                            "confidential",
                            "restricted",
                            "secret",
                            "level",
                        ]:
                            metadata[key] = value
        except Exception:
            pass
        return metadata

    @classmethod
    def _check_metadata_sensitivity(cls, metadata: dict[str, Any], blocked_levels: set) -> Optional[str]:
        """
        Check metadata for sensitivity indicators.

        Returns the detected sensitivity level if it matches blocked levels.
        """
        # Check direct sensitivity/classification fields
        fields = ["sensitivity", "classification", "mip_label", "category"]
        for field in fields:
            if field in metadata:
                value = str(metadata[field]).lower()
                if any(blocked in value for blocked in blocked_levels):
                    return value
        # Check other metadata fields for sensitivity keywords
        sensitive_fields = ["subject", "comments", "keywords"]
        for field in sensitive_fields:
            if field in metadata:
                value = str(metadata[field]).lower()
                for blocked in blocked_levels:
                    if blocked in value:
                        return str(blocked)
        # Check custom properties for MIP labels or sensitivity indicators
        for key, value in metadata.items():
            value_str = str(value).lower()
            # Look for Microsoft Information Protection patterns
            if any(
                pattern in key
                for pattern in [
                    "msip_label",
                    "sensitivity",
                    "classification",
                    "confidential",
                ]
            ):
                for blocked in blocked_levels:
                    if blocked in value_str:
                        return str(blocked)
        return None
