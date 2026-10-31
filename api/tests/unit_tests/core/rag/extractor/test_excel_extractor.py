import builtins
import importlib.util
import io
from collections.abc import Mapping, Sequence
from pathlib import Path
from types import ModuleType, SimpleNamespace

import pandas as pd
import pytest
from sqlalchemy import Engine, select
from sqlalchemy.orm import Session, sessionmaker

import core.rag.extractor.excel_extractor as excel_module
from core.rag.extractor.excel_extractor import ExcelExtractor
from models.base import TypeBase
from models.model import UploadFile
from tests.helpers.unit.config_override import apply_config_overrides


@pytest.fixture
def database_session_maker(sqlite_engine: Engine, monkeypatch: pytest.MonkeyPatch) -> sessionmaker[Session]:
    TypeBase.metadata.create_all(sqlite_engine, tables=[UploadFile.__table__])
    session_maker = sessionmaker(bind=sqlite_engine, expire_on_commit=False)
    monkeypatch.setattr(excel_module.session_factory, "create_session", session_maker)
    return session_maker


class _FakeCell:
    def __init__(self, value, hyperlink=None):
        self.value = value
        self.hyperlink = hyperlink
        self.row = 0
        self.column = 0


class _FakeSheet:
    def __init__(self, header_rows, data_rows, images=None):
        self._header_rows = header_rows
        self._data_rows = data_rows
        self._images = images or []
        self.max_row = len(header_rows) + len(data_rows)

    def iter_rows(self, min_row=1, max_row=None, max_col=None, values_only=False):
        if values_only:
            for row in self._header_rows:
                yield tuple(row)
            return

        for row_idx, row in enumerate(self._data_rows, start=min_row):
            materialized_row = tuple(row[:max_col] if max_col is not None else row)
            for col_idx, cell in enumerate(materialized_row, start=1):
                cell.row = row_idx
                cell.column = col_idx
            yield materialized_row


class _FakeWorkbook:
    def __init__(self, sheets):
        self._sheets = sheets
        self.sheetnames = list(sheets.keys())
        self.closed = False

    def __getitem__(self, key):
        return self._sheets[key]

    def close(self):
        self.closed = True


class _FakeImage:
    def __init__(self, data: bytes, row: int, col: int, image_format: str = "png"):
        self._raw_data = data
        self.anchor = SimpleNamespace(_from=SimpleNamespace(row=row, col=col))
        self.format = image_format

    def _data(self) -> bytes:
        return self._raw_data


def _patch_image_persistence(monkeypatch: pytest.MonkeyPatch):
    saves: list[tuple[str, bytes]] = []

    def save(key: str, data: bytes) -> None:
        saves.append((key, data))

    monkeypatch.setattr(excel_module.storage, "save", save)
    apply_config_overrides(monkeypatch, FILES_URL="http://files.local", STORAGE_TYPE="local")

    return saves


def _get_upload_files(session_maker: sessionmaker[Session]) -> list[UploadFile]:
    with session_maker() as session:
        return list(session.scalars(select(UploadFile)).all())


class TestExcelExtractor:
    def test_installed_xlsx_reader(self, tmp_path):
        import openpyxl

        if importlib.util.find_spec("wolfxl") is None:
            assert excel_module.load_workbook is openpyxl.load_workbook
        else:
            assert excel_module.load_workbook is importlib.import_module("wolfxl").load_workbook

        path = tmp_path / "reader.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append(["Name", "Count"])
        sheet.append(["Alice", 3])
        workbook.save(path)
        workbook.close()

        documents = ExcelExtractor(str(path)).extract()
        assert [document.page_content for document in documents] == ['"Name":"Alice";"Count":"3"']

    def test_broken_wolfxl_install_is_not_masked(self, monkeypatch: pytest.MonkeyPatch) -> None:
        real_import = builtins.__import__

        def import_with_broken_wolfxl(
            name: str,
            globals: Mapping[str, object] | None = None,
            locals: Mapping[str, object] | None = None,
            fromlist: Sequence[str] = (),
            level: int = 0,
        ) -> ModuleType:
            if name == "wolfxl":
                raise ModuleNotFoundError("No module named 'wolfxl._native'", name="wolfxl._native")
            return real_import(name, globals, locals, fromlist, level)

        monkeypatch.setattr(builtins, "__import__", import_with_broken_wolfxl)
        spec = importlib.util.spec_from_file_location("_excel_extractor_import_probe", excel_module.__file__)
        assert spec is not None
        assert spec.loader is not None

        with pytest.raises(ModuleNotFoundError) as exc_info:
            spec.loader.exec_module(importlib.util.module_from_spec(spec))
        assert exc_info.value.name == "wolfxl._native"

    def test_extract_xlsx_with_hyperlinks_and_sheet_skip(self, monkeypatch: pytest.MonkeyPatch):
        sheet_with_data = _FakeSheet(
            header_rows=[("Name", "Link")],
            data_rows=[
                (_FakeCell("Alice"), _FakeCell("Doc", hyperlink=SimpleNamespace(target="https://example.com/doc"))),
                (_FakeCell(None), _FakeCell(123)),
                (_FakeCell(None), _FakeCell(None)),
            ],
        )
        empty_sheet = _FakeSheet(header_rows=[(None, None)], data_rows=[])

        workbook = _FakeWorkbook({"Data": sheet_with_data, "Empty": empty_sheet})
        monkeypatch.setattr(excel_module, "load_workbook", lambda *args, **kwargs: workbook)

        extractor = ExcelExtractor("/tmp/sample.xlsx")
        docs = extractor.extract()

        assert workbook.closed is True
        assert len(docs) == 2
        assert docs[0].page_content == '"Name":"Alice";"Link":"[Doc](https://example.com/doc)"'
        assert docs[1].page_content == '"Name":"";"Link":"123"'
        assert all(doc.metadata["source"] == "/tmp/sample.xlsx" for doc in docs)

    def test_extract_xlsx_turns_embedded_images_into_markdown_links(
        self,
        monkeypatch: pytest.MonkeyPatch,
        database_session_maker: sessionmaker[Session],
    ):
        image_bytes = b"\x89PNG\r\n\x1a\nexcel-image"
        sheet = _FakeSheet(
            header_rows=[("Question", "Answer", "Image")],
            data_rows=[
                (_FakeCell("Q1"), _FakeCell("A1"), _FakeCell(None)),
                (_FakeCell("Q2"), _FakeCell("A2"), _FakeCell(None)),
            ],
            images=[
                _FakeImage(image_bytes, row=1, col=2),
                _FakeImage(image_bytes, row=1, col=2),
            ],
        )
        workbook = _FakeWorkbook({"Data": sheet})
        monkeypatch.setattr(excel_module, "load_workbook", lambda *args, **kwargs: workbook)
        saves = _patch_image_persistence(monkeypatch)

        extractor = ExcelExtractor(
            "/tmp/sample.xlsx",
            tenant_id="tenant-1",
            user_id="user-1",
            source_file_id="source-file-1",
        )
        docs = extractor.extract()
        upload_files = _get_upload_files(database_session_maker)

        assert workbook.closed is True
        assert len(docs) == 2
        assert len(upload_files) == 1
        assert docs[0].page_content == (
            '"Question":"Q1";"Answer":"A1";'
            f'"Image":"![image](http://files.local/files/{upload_files[0].id}/file-preview) '
            f'![image](http://files.local/files/{upload_files[0].id}/file-preview)"'
        )
        assert docs[1].page_content == '"Question":"Q2";"Answer":"A2";"Image":""'
        assert len(saves) == 1
        assert saves[0][0].startswith("image_files/tenant-1/source-file-1/")
        assert saves[0][0].endswith(".png")
        assert saves[0][1] == image_bytes
        assert upload_files[0].tenant_id == "tenant-1"
        assert upload_files[0].key == saves[0][0]
        assert upload_files[0].used is True

    def test_extract_xlsx_keeps_rows_with_only_embedded_images(
        self,
        monkeypatch: pytest.MonkeyPatch,
        database_session_maker: sessionmaker[Session],
    ):
        image_bytes = b"\x89PNG\r\n\x1a\nimage-only-row"
        sheet = _FakeSheet(
            header_rows=[("Question", "Answer", "Image")],
            data_rows=[
                (_FakeCell(None), _FakeCell(None), _FakeCell(None)),
                (_FakeCell(None), _FakeCell(None), _FakeCell(None)),
            ],
            images=[_FakeImage(image_bytes, row=1, col=2)],
        )
        workbook = _FakeWorkbook({"Data": sheet})
        monkeypatch.setattr(excel_module, "load_workbook", lambda *args, **kwargs: workbook)
        saves = _patch_image_persistence(monkeypatch)

        extractor = ExcelExtractor(
            "/tmp/sample.xlsx",
            tenant_id="tenant-1",
            user_id="user-1",
            source_file_id="source-file-1",
        )
        docs = extractor.extract()
        upload_files = _get_upload_files(database_session_maker)

        assert workbook.closed is True
        assert len(docs) == 1
        assert len(upload_files) == 1
        assert docs[0].page_content == (
            f'"Question":"";"Answer":"";"Image":"![image](http://files.local/files/{upload_files[0].id}/file-preview)"'
        )
        assert len(saves) == 1

    def test_extract_xlsx_reuses_existing_embedded_image_uploads_on_retry(
        self,
        monkeypatch: pytest.MonkeyPatch,
        database_session_maker: sessionmaker[Session],
    ):
        image_bytes = b"\x89PNG\r\n\x1a\nretry-safe-image"
        workbooks = [
            _FakeWorkbook(
                {
                    "Data": _FakeSheet(
                        header_rows=[("Question", "Answer", "Image")],
                        data_rows=[(_FakeCell("Q1"), _FakeCell("A1"), _FakeCell(None))],
                        images=[_FakeImage(image_bytes, row=1, col=2)],
                    )
                }
            ),
            _FakeWorkbook(
                {
                    "Data": _FakeSheet(
                        header_rows=[("Question", "Answer", "Image")],
                        data_rows=[(_FakeCell("Q1"), _FakeCell("A1"), _FakeCell(None))],
                        images=[_FakeImage(image_bytes, row=1, col=2)],
                    )
                }
            ),
        ]
        monkeypatch.setattr(excel_module, "load_workbook", lambda *args, **kwargs: workbooks.pop(0))
        saves = _patch_image_persistence(monkeypatch)

        extractor = ExcelExtractor(
            "/tmp/sample.xlsx",
            tenant_id="tenant-1",
            user_id="user-1",
            source_file_id="source-file-1",
        )
        first_docs = extractor.extract()
        second_docs = extractor.extract()
        upload_files = _get_upload_files(database_session_maker)
        assert len(upload_files) == 1

        expected_page_content = (
            '"Question":"Q1";"Answer":"A1";'
            f'"Image":"![image](http://files.local/files/{upload_files[0].id}/file-preview)"'
        )

        assert first_docs[0].page_content == expected_page_content
        assert second_docs[0].page_content == expected_page_content
        assert len(saves) == 1

    def test_extract_xlsx_keeps_trailing_image_only_rows_from_real_workbook(
        self,
        tmp_path: Path,
        monkeypatch: pytest.MonkeyPatch,
        database_session_maker: sessionmaker[Session],
    ) -> None:
        import openpyxl
        from openpyxl.drawing.image import Image as WorksheetImage
        from PIL import Image as PILImage

        png = io.BytesIO()
        PILImage.new("RGB", (4, 4), "red").save(png, format="PNG")
        image_bytes = png.getvalue()

        path = tmp_path / "images.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.append(["Question", "Answer", "Image"])
        sheet.append(["Q1", "A1", None])
        # Row 3 is just past the populated cells; row 12 is past the 10-row header scan.
        sheet.add_image(WorksheetImage(io.BytesIO(image_bytes)), "C3")
        sheet.add_image(WorksheetImage(io.BytesIO(image_bytes)), "C12")
        workbook.save(path)
        workbook.close()
        saves = _patch_image_persistence(monkeypatch)

        extractor = ExcelExtractor(
            str(path),
            tenant_id="tenant-1",
            user_id="user-1",
            source_file_id="source-file-1",
        )
        first_docs = extractor.extract()
        second_docs = extractor.extract()
        upload_files = _get_upload_files(database_session_maker)

        upload_ids = {
            row: next(upload_file.id for upload_file in upload_files if f"_r{row}_c3_" in upload_file.key)
            for row in (3, 12)
        }
        expected_page_contents = [
            '"Question":"Q1";"Answer":"A1";"Image":""',
            *(
                f'"Question":"";"Answer":"";"Image":"![image](http://files.local/files/{upload_ids[row]}/file-preview)"'
                for row in (3, 12)
            ),
        ]
        assert [doc.page_content for doc in first_docs] == expected_page_contents
        assert [doc.page_content for doc in second_docs] == expected_page_contents
        assert len(upload_files) == 2
        assert sorted(saves) == sorted((upload_file.key, image_bytes) for upload_file in upload_files)

    def test_extract_xls_path(self, monkeypatch: pytest.MonkeyPatch):
        class FakeExcelFile:
            sheet_names = ["Sheet1"]

            def parse(self, sheet_name):
                assert sheet_name == "Sheet1"
                return pd.DataFrame([{"A": "x", "B": 1}, {"A": None, "B": None}])

        monkeypatch.setattr(pd, "ExcelFile", lambda path, engine=None: FakeExcelFile())

        extractor = ExcelExtractor("/tmp/sample.xls")
        docs = extractor.extract()

        assert len(docs) == 1
        assert docs[0].page_content == '"A":"x";"B":"1.0"'
        assert docs[0].metadata == {"source": "/tmp/sample.xls"}

    def test_extract_xls_escapes_double_quotes_like_xlsx(self, monkeypatch: pytest.MonkeyPatch):
        class FakeExcelFile:
            sheet_names = ["Sheet1"]

            def parse(self, sheet_name):
                return pd.DataFrame([{"note": 'he said "hi"', "plain": "text"}])

        monkeypatch.setattr(pd, "ExcelFile", lambda path, engine=None: FakeExcelFile())

        extractor = ExcelExtractor("/tmp/sample.xls")
        docs = extractor.extract()

        assert len(docs) == 1
        assert docs[0].page_content == '"note":"he said \\"hi\\"";"plain":"text"'

    def test_extract_unsupported_extension_raises(self):
        extractor = ExcelExtractor("/tmp/sample.txt")

        with pytest.raises(ValueError, match="Unsupported file extension"):
            extractor.extract()

    def test_find_header_and_columns_prefers_first_row_with_two_columns(self):
        sheet = _FakeSheet(
            header_rows=[(None, None, None), ("A", "B", None), ("X", None, None)],
            data_rows=[],
        )
        extractor = ExcelExtractor("dummy.xlsx")

        header_row_idx, column_map, max_col_idx = extractor._find_header_and_columns(sheet)

        assert header_row_idx == 2
        assert column_map == {0: "A", 1: "B"}
        assert max_col_idx == 2

    def test_find_header_and_columns_fallback_and_empty_case(self):
        extractor = ExcelExtractor("dummy.xlsx")

        fallback_sheet = _FakeSheet(header_rows=[("Only", None), (None, "Second")], data_rows=[])
        row_idx, column_map, max_col_idx = extractor._find_header_and_columns(fallback_sheet)
        assert row_idx == 1
        assert column_map == {0: "Only"}
        assert max_col_idx == 1

        empty_sheet = _FakeSheet(header_rows=[(None, None)], data_rows=[])
        assert extractor._find_header_and_columns(empty_sheet) == (0, {}, 0)
