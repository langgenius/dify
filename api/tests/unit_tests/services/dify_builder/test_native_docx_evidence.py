"""Only owned bytes behind a native Tool output can prove a DOCX expectation."""

from collections.abc import Callable, Iterable, Iterator, Mapping
from io import BytesIO
from typing import NoReturn, TypedDict
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml.document import CT_Document

from core.dify_builder.models import Graph, NodeOutput, Run
from core.workflow.file_reference import build_file_reference
from services.dify_builder.native_docx_evidence import FileAcceptance, StoredToolFile, validate_native_docx
from tests.unit_tests.services.dify_builder.test_business_acceptance import GRAPH, observation

FILE_ID = "b50c9ac2-93f9-4bec-82e4-ac4068fb2bf8"
TENANT = "tenant"
KEY = "tools/tenant/report.docx"
CONTENT_TYPES = (
    b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    b'<Override PartName="/word/document.xml" '
    b'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    b"</Types>"
)
DOCUMENT = (
    b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    b"<w:body><w:p/></w:body></w:document>"
)


ROOT_RELATIONSHIPS = (
    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    b'<Relationship Id="rId1" '
    b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    b'Target="word/document.xml"/></Relationships>'
)


def docx(parts: Mapping[str, bytes] | None = None) -> bytes:
    buf = BytesIO()
    if parts is None:
        document = Document()
        document.add_paragraph("Native report")
        document.save(buf)
    else:
        with ZipFile(buf, "w", ZIP_DEFLATED) as package:
            for name, value in parts.items():
                package.writestr(name, value)
    return buf.getvalue()


def native_file(**changes: object) -> dict[str, object]:
    return {
        "dify_model_identity": "__dify__file__",
        "transfer_method": "tool_file",
        "type": "document",
        "reference": build_file_reference(record_id=FILE_ID),
        "related_id": FILE_ID,
        "filename": "report.docx",
        "extension": ".docx",
        "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        **changes,
    }


class FileRowChanges(TypedDict, total=False):
    id: str
    tenant_id: str
    file_key: str
    size: int


def assess(
    value: object = None,
    blob: bytes | None = None,
    row_changes: FileRowChanges | None = None,
    *,
    run: Run | None = None,
    native_graph: Graph | None = None,
    lookup_file: Callable[[str, str], StoredToolFile | None] | None = None,
    load_stream: Callable[[str], Iterable[bytes]] | None = None,
    max_file_bytes: int = 10 * 1024 * 1024,
    max_uncompressed_bytes: int = 20 * 1024 * 1024,
) -> tuple[FileAcceptance, list[tuple[str, ...]]]:
    value = native_file() if value is None else value
    blob = docx() if blob is None else blob
    observed_run, _ = observation({"files": [value]})
    observed_run.per_node[0] = NodeOutput(
        node_id="json", status="succeeded", outputs={"files": [native_file()]}, outputs_available=True
    )
    calls: list[tuple[str, ...]] = []
    row_changes = row_changes or {}

    def lookup(tenant_id: str, file_id: str) -> StoredToolFile:
        calls.append(("lookup", tenant_id, file_id))
        return StoredToolFile(
            id=row_changes.get("id", FILE_ID),
            tenant_id=row_changes.get("tenant_id", TENANT),
            file_key=row_changes.get("file_key", KEY),
            size=row_changes.get("size", len(blob)),
        )

    def read(key: str) -> Iterator[bytes]:
        calls.append(("read", key))
        yield blob

    result = validate_native_docx(
        value,
        tenant_id=TENANT,
        run=observed_run if run is None else run,
        native_graph=GRAPH if native_graph is None else native_graph,
        lookup_file=lookup if lookup_file is None else lookup_file,
        load_stream=read if load_stream is None else load_stream,
        max_file_bytes=max_file_bytes,
        max_uncompressed_bytes=max_uncompressed_bytes,
    )
    return result, calls


def test_actual_small_docx_zip_from_native_tool_file_passes() -> None:
    result, calls = assess()
    assert result.state == "passed"
    assert result.reason == "docx_bytes_validated"
    assert Document(BytesIO(docx())).paragraphs[0].text == "Native report"
    assert calls == [("lookup", TENANT, FILE_ID), ("read", KEY)]


@pytest.mark.parametrize(
    "value",
    [
        "report.docx",
        {"filename": "report.docx", "url": "https://example.com/report.docx"},
        native_file(transfer_method="remote_url"),
        native_file(related_id="another-file"),
        native_file(reference="not-a-file-id"),
        native_file(dify_model_identity="fake"),
        native_file(reference="dify-file-ref:W10="),
    ],
)
def test_filename_url_remote_or_conflicting_identity_never_reads_storage(value: object) -> None:
    result, calls = assess(value)
    assert result.state == "unverified"
    assert calls == []


def test_wrong_tenant_row_never_reads_storage() -> None:
    result, calls = assess(row_changes={"tenant_id": "foreign"})
    assert result.state == "unverified"
    assert calls == [("lookup", TENANT, FILE_ID)]


def test_storage_key_is_taken_only_from_owned_database_row() -> None:
    result, calls = assess(native_file(reference=build_file_reference(record_id=FILE_ID, storage_key="other/key")))
    assert result.state == "unverified"
    assert calls == [("lookup", TENANT, FILE_ID)]


def test_unrelated_or_failed_tool_output_cannot_prove_terminal_file() -> None:
    run, _ = observation({"files": [native_file()]})
    run.per_node[0] = NodeOutput(
        node_id="json", status="failed", outputs={"files": [native_file()]}, outputs_available=True
    )
    result, calls = assess(run=run)
    assert result.state == "unverified"
    assert calls == []


@pytest.mark.parametrize(
    "blob",
    [
        b"renamed text",
        b"PK malformed",
        docx({"word/document.xml": DOCUMENT}),
        docx({"[Content_Types].xml": b"not XML", "word/document.xml": DOCUMENT, "_rels/.rels": ROOT_RELATIONSHIPS}),
        docx(
            {"[Content_Types].xml": CONTENT_TYPES, "word/document.xml": b"<fake/>", "_rels/.rels": ROOT_RELATIONSHIPS}
        ),
    ],
)
def test_readable_invalid_docx_bytes_fail_even_with_docx_filename(blob: bytes) -> None:
    result, _ = assess(blob=blob)
    assert result.state == "failed"


def test_declared_size_limit_prevents_storage_read() -> None:
    result, calls = assess(row_changes={"size": 99999}, max_file_bytes=1000)
    assert result.state == "unverified"
    assert calls == [("lookup", TENANT, FILE_ID)]


def test_stream_is_bounded_and_closed_when_actual_bytes_exceed_limit() -> None:
    events: list[str] = []

    def read(_key: str) -> Iterator[bytes]:
        try:
            events.append("first")
            yield b"x" * 1001
            events.append("second")
            yield b"not read"
        finally:
            events.append("closed")

    result, _ = assess(row_changes={"size": -1}, max_file_bytes=1000, load_stream=read)
    assert result.state == "unverified"
    assert events == ["first", "closed"]


def test_plain_iterable_stream_needs_no_close_method() -> None:
    blob = docx()
    result, _ = assess(blob=blob, load_stream=lambda _key: [blob])
    assert result.state == "passed"


def test_non_generator_closeable_stream_releases_owner_on_completion() -> None:
    events: list[str] = []
    blob = docx()

    class OwnedByteStream:
        def __iter__(self) -> Iterator[bytes]:
            yield blob

        def close(self) -> None:
            events.append("closed")

    result, _ = assess(blob=blob, load_stream=lambda _key: OwnedByteStream())
    assert result.state == "passed"
    assert events == ["closed"]


def test_zip_uncompressed_size_bound_is_checked_before_expansion() -> None:
    result, _ = assess(
        blob=docx({"[Content_Types].xml": CONTENT_TYPES, "word/document.xml": DOCUMENT, "large": b"x" * 20000}),
        max_uncompressed_bytes=1000,
    )
    assert result.state == "unverified"


@pytest.mark.parametrize("owner", ["lookup", "read"])
def test_unavailable_database_or_storage_bytes_remain_unverified(owner: str) -> None:
    def unavailable(*_args: object) -> NoReturn:
        raise OSError("unavailable")

    result, _ = assess(lookup_file=unavailable) if owner == "lookup" else assess(load_stream=unavailable)
    assert result.state == "unverified"


def test_file_list_passes_only_when_a_real_docx_was_validated() -> None:
    result, _ = assess([native_file()])
    assert result.state == "passed"
    result, calls = assess([])
    assert result.state == "failed"
    assert calls == []


def test_actual_sql_lookup_scopes_tool_file_to_tenant(monkeypatch: pytest.MonkeyPatch) -> None:
    from sqlalchemy import Table, create_engine
    from sqlalchemy.orm import sessionmaker

    from core.db.session_factory import session_factory
    from models.tools import ToolFile
    from services.dify_builder.native_docx_evidence import lookup_owned_tool_file

    engine = create_engine("sqlite://")
    table = ToolFile.__table__
    assert isinstance(table, Table)
    table.create(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    owner = "cff96ac1-a314-4b94-80da-06c43ea0365e"
    with factory() as sql:
        row = ToolFile(
            user_id="899bb86c-ddb7-4cb9-a064-5e710ce3c918",
            tenant_id=owner,
            conversation_id=None,
            file_key=f"tools/{owner}/report.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            size=123,
        )
        row.id = FILE_ID
        sql.add(row)
        sql.commit()
    monkeypatch.setattr(session_factory, "create_session", factory)
    assert lookup_owned_tool_file("7a108fa2-022d-41c7-bfd5-5d93d354c9e6", FILE_ID) is None
    row = lookup_owned_tool_file(owner, FILE_ID)
    assert row is not None
    assert row.tenant_id == owner
    assert row.file_key == f"tools/{owner}/report.docx"
    engine.dispose()


@pytest.mark.parametrize("owner", ["lookup", "read"])
def test_file_annotation_lookup_or_read_error_preserves_completed_run(
    monkeypatch: pytest.MonkeyPatch, owner: str
) -> None:
    from core.dify_builder.models import DifyBuilderContext, EntryMode, Session
    from core.dify_builder.state import PcState
    from extensions.ext_storage import storage
    from services.dify_builder import native_docx_evidence
    from services.dify_builder.business_acceptance import confirm_testdata_contract, record_run_acceptance
    from tests.unit_tests.core.dify_builder.fakes import InMemoryRepository

    value = native_file()
    run, ti = observation({"document": [value]})
    run.per_node[0] = NodeOutput(node_id="json", status="succeeded", outputs={"files": [value]}, outputs_available=True)
    s = Session(id="s", app_id="app", tenant_id=TENANT, entry_mode=EntryMode.FIX, current_state=PcState.FIX_VERIFY)
    context = DifyBuilderContext(goal_text="Explicit file expectation", verify_run_id=run.id)
    confirm_testdata_contract(
        s,
        context,
        {
            "business_acceptance_contract": {
                "confirmed": True,
                "cases": [
                    {
                        "id": "paid",
                        "inputs": ti.inputs,
                        "checks": [{"path": ["document"], "operator": "file", "file_type": "docx"}],
                    }
                ],
            },
            "acceptance_case_id": "paid",
        },
        ti.inputs,
    )

    def unavailable(*_args: object) -> NoReturn:
        raise OSError("Storage owner unavailable")

    if owner == "lookup":
        monkeypatch.setattr(native_docx_evidence, "lookup_owned_tool_file", unavailable)
    else:
        monkeypatch.setattr(
            native_docx_evidence,
            "lookup_owned_tool_file",
            lambda _tenant, _id: native_docx_evidence.StoredToolFile(FILE_ID, TENANT, KEY, -1),
        )
        monkeypatch.setattr(storage, "load_stream", unavailable)
    record_run_acceptance(s, context, run, ti, GRAPH, "rev")
    assert run.status == "succeeded"
    assert run.verification is not None
    assert run.verification.business_acceptance["state"] == "unverified"
    assert context.verify_run_id == run.id
    repo = InMemoryRepository()
    repo.save_run(s.id, run)
    saved_run = repo.get_run(run.id)
    assert saved_run is not None
    assert saved_run.verification is not None
    assert saved_run.verification.terminal_outputs == {"document": [value]}


def test_file_lookup_requires_the_graph_bound_to_immutable_native_run() -> None:
    from copy import deepcopy

    changed: Graph = deepcopy(GRAPH)
    changed["nodes"][0]["data"]["tool_name"] = "another-tool"
    result, calls = assess(native_graph=changed)
    assert result.state == "unverified"
    assert calls == []


def test_missing_root_document_relationship_cannot_prove_docx() -> None:
    blob = docx({"[Content_Types].xml": CONTENT_TYPES, "word/document.xml": DOCUMENT})
    result, _ = assess(blob=blob)
    assert result.state == "failed"


@pytest.mark.parametrize(
    "relationships",
    [
        b"not XML",
        ROOT_RELATIONSHIPS.replace(b' Id="rId1"', b""),
        ROOT_RELATIONSHIPS.replace(b"rId1", b""),
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>',
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        b'Target="https://example.com/report.docx" TargetMode="External"/></Relationships>',
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        b'Target="../word/document.xml"/></Relationships>',
    ],
)
def test_invalid_or_external_root_relationship_cannot_prove_docx(relationships: bytes) -> None:
    blob = docx({"[Content_Types].xml": CONTENT_TYPES, "word/document.xml": DOCUMENT, "_rels/.rels": relationships})
    assert assess(blob=blob)[0].state == "failed"


@pytest.mark.parametrize("target", [b"missing.xml", b"%2e%2e/word/document.xml", b"word/document.xml#fragment"])
def test_root_relationship_requires_existing_safe_package_part(target: bytes) -> None:
    blob = docx(
        {
            "[Content_Types].xml": CONTENT_TYPES,
            "word/document.xml": DOCUMENT,
            "_rels/.rels": ROOT_RELATIONSHIPS.replace(b"word/document.xml", target),
        }
    )
    assert assess(blob=blob)[0].state == "failed"


def test_duplicate_office_document_relationship_is_not_a_verified_docx() -> None:
    extra = ROOT_RELATIONSHIPS.split(b">", 1)[1].removesuffix(b"</Relationships>").replace(b"rId1", b"rId2")
    relationships = ROOT_RELATIONSHIPS.replace(b"</Relationships>", extra + b"</Relationships>")
    blob = docx({"[Content_Types].xml": CONTENT_TYPES, "word/document.xml": DOCUMENT, "_rels/.rels": relationships})
    assert assess(blob=blob)[0].state == "failed"


def test_main_part_must_have_the_docx_content_type_declaration() -> None:
    blob = docx(
        {
            "[Content_Types].xml": CONTENT_TYPES.replace(b"document.main+xml", b"template.main+xml"),
            "word/document.xml": DOCUMENT,
            "_rels/.rels": ROOT_RELATIONSHIPS,
        }
    )
    assert assess(blob=blob)[0].state == "failed"


def test_internal_root_relationship_resolves_a_declared_readable_main_part() -> None:
    blob = docx(
        {
            "[Content_Types].xml": CONTENT_TYPES.replace(b"/word/document.xml", b"/report.xml"),
            "report.xml": DOCUMENT,
            "_rels/.rels": ROOT_RELATIONSHIPS.replace(b"word/document.xml", b"/report.xml"),
        }
    )
    element = Document(BytesIO(blob)).element
    assert isinstance(element, CT_Document)
    assert element.body is not None
    assert assess(blob=blob)[0].state == "passed"


def test_root_relationship_xml_has_the_same_read_bound_as_document_xml() -> None:
    blob = docx(
        {
            "[Content_Types].xml": CONTENT_TYPES,
            "word/document.xml": DOCUMENT,
            "_rels/.rels": ROOT_RELATIONSHIPS + b" " * (2 * 1024 * 1024),
        }
    )
    result, _ = assess(blob=blob)
    assert result.state == "unverified"
    assert result.reason == "docx_xml_limit"


def test_conflicting_main_part_content_types_do_not_prove_docx() -> None:
    duplicate = b'<Override PartName="/word/document.xml" ContentType="text/plain"/>'
    blob = docx(
        {
            "[Content_Types].xml": CONTENT_TYPES.replace(b"</Types>", duplicate + b"</Types>"),
            "word/document.xml": DOCUMENT,
            "_rels/.rels": ROOT_RELATIONSHIPS,
        }
    )
    assert assess(blob=blob)[0].state == "failed"
