"""Exercise the opt-in local context profile with actual file bytes."""

import io
import zipfile
from collections.abc import Iterator

import pytest

from services.dify_builder.attachment_extraction import ExtractedContext


def extract(extension: str, data: bytes) -> ExtractedContext:
    from services.dify_builder.attachment_extraction import extract_local_context

    return extract_local_context(extension, iter([data]))


def test_text_is_bounded_and_records_its_exact_digest() -> None:
    import hashlib

    value = extract("txt", b"a" * 20001)
    assert value.status == "truncated"
    assert value.text == "a" * 20000
    assert value.digest == hashlib.sha256(b"a" * 20001).hexdigest()


def test_native_stream_stops_before_materializing_an_oversized_file() -> None:
    from services.dify_builder.attachment_extraction import extract_local_context

    def chunks() -> Iterator[bytes]:
        yield b"a" * (10 * 1024 * 1024 + 1)
        raise AssertionError("must not continue an oversized stream")

    assert extract_local_context("txt", chunks()).reason == "byte_limit"


@pytest.mark.parametrize(
    ("extension", "data", "reason"),
    [
        ("png", b"not an image", "unsupported_format"),
        ("pdf", b"not a PDF", "malformed_file"),
        ("txt", b"\xff\xfe\x00", "unsupported_encoding"),
    ],
)
def test_finite_rejection_statuses(extension: str, data: bytes, reason: str) -> None:
    result = extract(extension, data)
    assert result.status == "rejected"
    assert result.reason == reason
    assert result.text == ""


def test_xlsx_reads_values_without_hyperlinks_or_formulas() -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["Invoice", 4, '=HYPERLINK("https://unsafe.test","Open")'])
    blob = io.BytesIO()
    book.save(blob)
    result = extract("xlsx", blob.getvalue())
    assert result.status == "ready"
    assert "Invoice\t4" in result.text
    assert "unsafe.test" not in result.text
    assert "HYPERLINK" not in result.text


def test_xlsx_rejects_zip_expansion_before_loading_workbook() -> None:
    blob = io.BytesIO()
    with zipfile.ZipFile(blob, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", "a" * (21 * 1024 * 1024))
    assert extract("xlsx", blob.getvalue()).reason == "zip_expansion_limit"


def test_pdf_page_budget_and_image_only_are_explicit() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    blob = io.BytesIO()
    writer.write(blob)
    assert extract("pdf", blob.getvalue()).reason == "no_extractable_text"
    for _ in range(100):
        writer.add_blank_page(width=72, height=72)
    blob = io.BytesIO()
    writer.write(blob)
    assert extract("pdf", blob.getvalue()).reason == "page_limit"


def test_native_upload_byte_limit_is_enforced_against_actual_bytes() -> None:
    from services.dify_builder.attachment_extraction import extract_local_context

    assert extract_local_context("txt", [b"abc"], max_bytes=2).reason == "byte_limit"


@pytest.mark.parametrize(("sheets", "cells", "reason"), [(11, 1, "sheet_limit"), (1, 10001, "cell_limit")])
def test_xlsx_structural_budgets(sheets: int, cells: int, reason: str) -> None:
    from openpyxl import Workbook

    book = Workbook()
    for index in range(sheets - 1):
        book.create_sheet(str(index))
    sheet = book.active
    assert sheet is not None
    for _ in range(cells):
        sheet.append([1])
    blob = io.BytesIO()
    book.save(blob)
    assert extract("xlsx", blob.getvalue()).reason == reason


def test_pdf_rejects_expanded_page_content_before_text_interpretation() -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=72, height=72)
    content = DecodedStreamObject()
    content.set_data(b" " * (2 * 1024 * 1024 + 1))
    page[NameObject("/Contents")] = writer._add_object(content.flate_encode())
    blob = io.BytesIO()
    writer.write(blob)
    assert extract("pdf", blob.getvalue()).reason == "pdf_content_limit"


def test_encrypted_pdf_has_an_explicit_local_rejection() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.encrypt("password")
    blob = io.BytesIO()
    writer.write(blob)
    assert extract("pdf", blob.getvalue()).reason == "encrypted_file"


def test_encrypted_office_container_is_not_misreported_as_supported_xlsx() -> None:
    assert extract("xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"opaque").reason == "encrypted_file"


def test_pdf_form_content_is_honestly_unsupported_before_nested_decoding() -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=72, height=72)
    form = DecodedStreamObject()
    form.set_data(b" " * (2 * 1024 * 1024 + 1))
    form[NameObject("/Subtype")] = NameObject("/Form")
    form_ref = writer._add_object(form.flate_encode())
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/XObject"): DictionaryObject({NameObject("/Fm"): form_ref})}
    )
    contents = DecodedStreamObject()
    contents.set_data(b"q /Fm Do Q")
    page[NameObject("/Contents")] = writer._add_object(contents)
    blob = io.BytesIO()
    writer.write(blob)
    assert extract("pdf", blob.getvalue()).reason == "unsupported_pdf_content"


def test_pdf_image_decoding_filter_cannot_enter_the_local_text_pipeline() -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=72, height=72)
    contents = DecodedStreamObject()
    contents.set_data(b"opaque")
    contents[NameObject("/Filter")] = NameObject("/JBIG2Decode")
    page[NameObject("/Contents")] = writer._add_object(contents)
    blob = io.BytesIO()
    writer.write(blob)
    assert extract("pdf", blob.getvalue()).reason == "unsupported_pdf_filter"
