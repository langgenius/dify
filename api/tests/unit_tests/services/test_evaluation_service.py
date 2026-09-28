"""Focused tests for evaluation service validation and dataset parsing."""

import io
import json
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from core.evaluation.entities.config_entity import EvaluationFrameworkEnum
from core.evaluation.entities.evaluation_entity import (
    AgentToolCall,
    CustomizedMetricOutputField,
    CustomizedMetrics,
    DefaultMetric,
    EvaluationCategory,
    EvaluationConfigData,
    EvaluationDatasetInput,
    EvaluationRunRequest,
    EvaluationRunSnapshot,
    NodeInfo,
)
from core.evaluation.entities.judgment_entity import JudgmentCondition, JudgmentConfig
from core.evaluation.evaluation_manager import EvaluationManager
from core.evaluation.exceptions import EvaluationWorkflowVersionNotFoundError
from libs.datetime_utils import naive_utc_now
from models.account import Tenant
from models.evaluation import (
    EvaluationConfiguration,
    EvaluationRun,
    EvaluationRunItem,
    EvaluationRunItemStatus,
    EvaluationRunStatus,
)
from models.model import App, AppMode
from models.workflow import WorkflowKind, WorkflowType
from services.errors.evaluation import EvaluationDatasetInvalidError
from services.evaluation_service import EvaluationService
from tests.unit_tests.model_factories import make_upload_file


@pytest.fixture(autouse=True)
def stub_evaluation_model_validation(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(EvaluationService, "_validate_evaluation_model", MagicMock())


def test_generated_template_uses_provided_column_names() -> None:
    content = EvaluationService._generate_xlsx_template(
        [
            {"name": "index", "type": "number"},
            {"name": "customer_question", "type": "string"},
        ],
        "App",
    )
    worksheet = load_workbook(io.BytesIO(content)).active

    assert worksheet is not None
    assert [cell.value for cell in worksheet[1]] == ["index", "customer_question"]


def test_generated_template_includes_reference_columns_for_selected_metrics(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    app = App(name="Reference App", mode=AppMode.WORKFLOW.value)
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_expected_output.return_value = True
    evaluation_instance.metric_requires_expected_tool_calls.return_value = False
    monkeypatch.setattr(
        EvaluationService,
        "_get_app_input_fields",
        MagicMock(return_value=[{"variable": "input", "label": "Input label", "type": "text-input"}]),
    )
    monkeypatch.setattr(
        EvaluationManager,
        "get_evaluation_instance",
        MagicMock(return_value=evaluation_instance),
    )

    content, _ = EvaluationService.generate_dataset_template(
        app,
        "apps",
        EvaluationConfigData(
            default_metrics=[
                DefaultMetric(
                    metric="answer_correctness",
                    node_info_list=[NodeInfo(node_id="llm-1", type="llm", title="Answer")],
                )
            ]
        ),
        MagicMock(spec=Session),
    )
    worksheet = load_workbook(io.BytesIO(content)).active

    assert worksheet is not None
    assert [cell.value for cell in worksheet[1]] == ["index", "input", "Answer : expected_output"]


def test_parse_csv_dataset_preserves_node_title_scoped_references() -> None:
    items = EvaluationService._parse_dataset(
        b"index,query,Answer : expected_output\n7,hello,world\n",
        "dataset.csv",
    )

    assert len(items) == 1
    assert items[0].index == 7
    assert items[0].inputs == {"query": "hello"}
    assert items[0].expected_outputs == {"Answer": "world"}


def test_parse_dataset_preserves_node_title_scoped_expected_tool_calls() -> None:
    items = EvaluationService._rows_to_dataset_items(
        [
            ["index", "query", "Research Agent : expected_tool_calls"],
            [
                7,
                "weather in Paris",
                json.dumps([{"name": "weather", "arguments": {"city": "Paris"}}]),
            ],
        ]
    )

    assert items[0].inputs == {"query": "weather in Paris"}
    assert items[0].expected_tool_calls_by_node == {
        "Research Agent": [AgentToolCall(name="weather", arguments={"city": "Paris"})]
    }
    assert items[0].iter_expected_output_columns() == [
        (
            "Research Agent : expected_tool_calls",
            '[{"arguments": {"city": "Paris"}, "name": "weather"}]',
        )
    ]


def test_parse_dataset_rejects_invalid_expected_tool_calls() -> None:
    with pytest.raises(EvaluationDatasetInvalidError, match="JSON array of tool calls.*7"):
        EvaluationService._rows_to_dataset_items(
            [
                ["index", "query", "Research Agent : expected_tool_calls"],
                [7, "hello", '{"name": "weather"}'],
            ]
        )


def test_parse_csv_dataset_rejects_duplicate_indices() -> None:
    with pytest.raises(EvaluationDatasetInvalidError, match="index '1' is duplicated"):
        EvaluationService._parse_dataset(
            b"index,query\n1,first\n1,second\n",
            "dataset.csv",
        )


def test_parse_csv_dataset_preserves_supported_scalar_input_types() -> None:
    items = EvaluationService._parse_dataset(
        b"index,query,temperature,enabled\n1,hello,2,true\n2,world,0.75,0\n3,large,9007199254740993,YES\n",
        "dataset.csv",
        input_types={"query": "string", "temperature": "number", "enabled": "boolean"},
    )

    assert items[0].inputs == {"query": "hello", "temperature": 2, "enabled": True}
    assert items[1].inputs == {"query": "world", "temperature": 0.75, "enabled": False}
    assert items[2].inputs == {"query": "large", "temperature": 9007199254740993, "enabled": True}


def test_parse_csv_dataset_rejects_invalid_number_input() -> None:
    with pytest.raises(EvaluationDatasetInvalidError, match="temperature.*number.*7"):
        EvaluationService._parse_dataset(
            b"index,temperature\n7,not-a-number\n",
            "dataset.csv",
            input_types={"temperature": "number"},
        )


def test_parse_csv_dataset_rejects_invalid_boolean_input() -> None:
    with pytest.raises(EvaluationDatasetInvalidError, match="enabled.*boolean.*8"):
        EvaluationService._parse_dataset(
            b"index,enabled\n8,maybe\n",
            "dataset.csv",
            input_types={"enabled": "boolean"},
        )


def test_dataset_input_columns_reject_unsupported_types() -> None:
    assert EvaluationService._dataset_input_types([{"variable": "enabled", "type": "checkbox"}]) == {
        "enabled": "boolean"
    }
    with pytest.raises(ValueError, match="only support string, number, and boolean.*file"):
        EvaluationService._dataset_input_types([{"variable": "attachment", "type": "file"}])


def test_dataset_validation_requires_references_for_reference_metrics() -> None:
    items = [EvaluationDatasetInput(index=7, inputs={"query": "hello"})]
    metrics = [DefaultMetric(metric="context_recall", node_info_list=[])]

    with pytest.raises(EvaluationDatasetInvalidError, match="expected_output.*7"):
        EvaluationService._validate_dataset_expected_outputs(items, metrics)


def test_dataset_validation_allows_reference_free_metrics() -> None:
    items = [EvaluationDatasetInput(index=7, inputs={"query": "hello"})]
    metrics = [DefaultMetric(metric="context_relevance", node_info_list=[])]

    EvaluationService._validate_dataset_expected_outputs(items, metrics)


def test_dataset_validation_resolves_references_by_node_title() -> None:
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_expected_output.return_value = True
    metrics = [
        DefaultMetric(
            metric="answer_correctness",
            node_info_list=[
                NodeInfo(node_id="llm-1", type="llm", title="Primary Answer"),
                NodeInfo(node_id="llm-2", type="llm", title="Backup Answer"),
            ],
        )
    ]

    EvaluationService._validate_dataset_expected_outputs(
        [
            EvaluationDatasetInput(
                index=7,
                inputs={"query": "hello"},
                expected_outputs={"Primary Answer": "first", "Backup Answer": "second"},
            )
        ],
        metrics,
        evaluation_instance=evaluation_instance,
    )

    with pytest.raises(EvaluationDatasetInvalidError, match="expected_output.*7"):
        EvaluationService._validate_dataset_expected_outputs(
            [
                EvaluationDatasetInput(
                    index=7,
                    inputs={"query": "hello"},
                    expected_outputs={"Primary Answer": "first"},
                )
            ],
            metrics,
            evaluation_instance=evaluation_instance,
        )


def test_dataset_validation_requires_reference_tool_calls_for_ragas_agent_metric() -> None:
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_expected_output.return_value = False
    evaluation_instance.metric_requires_expected_tool_calls.side_effect = lambda metric: metric == "tool_call_accuracy"
    metrics = [
        DefaultMetric(
            metric="tool_call_accuracy",
            node_info_list=[NodeInfo(node_id="agent-1", type="agent", title="Research Agent")],
        )
    ]

    with pytest.raises(EvaluationDatasetInvalidError, match="expected_tool_calls.*7"):
        EvaluationService._validate_dataset_expected_outputs(
            [EvaluationDatasetInput(index=7, inputs={"query": "hello"})],
            metrics,
            evaluation_instance=evaluation_instance,
        )

    with pytest.raises(EvaluationDatasetInvalidError, match="expected_tool_calls.*7"):
        EvaluationService._validate_dataset_expected_outputs(
            [
                EvaluationDatasetInput(
                    index=7,
                    inputs={"query": "hello"},
                    expected_tool_calls_by_node={
                        "Other Agent": [AgentToolCall(name="search", arguments={"query": "hello"})]
                    },
                )
            ],
            metrics,
            evaluation_instance=evaluation_instance,
        )

    EvaluationService._validate_dataset_expected_outputs(
        [
            EvaluationDatasetInput(
                index=7,
                inputs={"query": "hello"},
                expected_tool_calls_by_node={
                    "Research Agent": [AgentToolCall(name="search", arguments={"query": "hello"})]
                },
            )
        ],
        metrics,
        evaluation_instance=evaluation_instance,
    )


def test_expected_tool_call_columns_are_generated_for_reference_based_agent_metrics() -> None:
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_expected_tool_calls.side_effect = lambda metric: metric == "tool_call_accuracy"

    columns = EvaluationService._get_expected_tool_call_columns(
        [
            DefaultMetric(
                metric="tool_call_accuracy",
                node_info_list=[
                    NodeInfo(node_id="agent-1", type="agent", title="Research Agent"),
                    NodeInfo(node_id="agent-2", type="agent", title="Backup Agent"),
                ],
            ),
            DefaultMetric(
                metric="argument_correctness",
                node_info_list=[NodeInfo(node_id="agent-1", type="agent", title="Research Agent")],
            ),
        ],
        evaluation_instance=evaluation_instance,
    )

    assert columns == [
        {"name": "Research Agent : expected_tool_calls", "type": "string"},
        {"name": "Backup Agent : expected_tool_calls", "type": "string"},
    ]


def test_expected_output_columns_use_node_titles() -> None:
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_expected_output.return_value = True

    columns = EvaluationService._get_expected_output_columns(
        [
            DefaultMetric(
                metric="answer_correctness",
                node_info_list=[
                    NodeInfo(node_id="llm-1", type="llm", title="Primary Answer"),
                    NodeInfo(node_id="llm-2", type="llm", title="Backup Answer"),
                ],
            )
        ],
        evaluation_instance=evaluation_instance,
    )

    assert columns == [
        {"name": "Primary Answer : expected_output", "type": "string"},
        {"name": "Backup Answer : expected_output", "type": "string"},
    ]


def test_dataset_references_fall_back_to_legacy_node_ids() -> None:
    expected_tool_calls = [AgentToolCall(name="search", arguments={"query": "hello"})]
    item = EvaluationDatasetInput(
        index=7,
        inputs={"query": "hello"},
        expected_outputs={"llm-1": "legacy answer"},
        expected_tool_calls_by_node={"agent-1": expected_tool_calls},
    )

    assert item.get_expected_output_for_node("llm-1", "Answer") == "legacy answer"
    assert item.get_expected_tool_calls_for_node("agent-1", "Research Agent") == expected_tool_calls


def test_validate_config_rejects_node_type_spoofing(monkeypatch: pytest.MonkeyPatch) -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    workflow = SimpleNamespace(walk_nodes=lambda: iter([("node-1", {"type": "llm"})]))
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=workflow))
    monkeypatch.setattr(
        EvaluationService,
        "get_supported_metrics",
        MagicMock(return_value=["context_precision"]),
    )
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[
            DefaultMetric(
                metric="context_precision",
                node_info_list=[NodeInfo(node_id="node-1", type="knowledge-retrieval", title="Retriever")],
            )
        ],
    )

    with pytest.raises(ValueError, match="type does not match"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="apps",
            target_id="app-1",
            data=config,
        )


def test_validate_config_rejects_agent_v2(monkeypatch: pytest.MonkeyPatch) -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    workflow = SimpleNamespace(
        walk_nodes=lambda: iter(
            [
                (
                    "agent-1",
                    {"type": "agent", "version": "2", "agent_node_kind": "dify_agent"},
                )
            ]
        ),
        user_input_form=lambda: [],
    )
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=workflow))
    monkeypatch.setattr(EvaluationService, "get_supported_metrics", MagicMock(return_value=["argument_correctness"]))
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_embeddings.return_value = False
    monkeypatch.setattr(
        "services.evaluation_service.EvaluationManager.get_evaluation_instance",
        MagicMock(return_value=evaluation_instance),
    )
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[
            DefaultMetric(
                metric="argument_correctness",
                node_info_list=[NodeInfo(node_id="agent-1", type="agent", title="Agent")],
            )
        ],
    )

    with pytest.raises(ValueError, match="Agent V2.*does not support"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="apps",
            target_id="app-1",
            data=config,
        )


def test_agent_node_discovery_includes_legacy_agent_and_excludes_agent_v2() -> None:
    workflow = SimpleNamespace(
        walk_nodes=lambda: iter(
            [
                ("legacy-agent", {"type": "agent", "version": "2", "title": "Legacy Agent"}),
                (
                    "agent-v2",
                    {
                        "type": "agent",
                        "version": "2",
                        "agent_node_kind": "dify_agent",
                        "title": "Agent V2",
                    },
                ),
            ]
        )
    )

    nodes = EvaluationService._nodes_for_metrics_from_workflow(workflow, ["argument_correctness"])

    assert nodes == {"argument_correctness": [{"node_id": "legacy-agent", "type": "agent", "title": "Legacy Agent"}]}


def test_validate_config_requires_default_embedding_for_embedding_metrics(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    workflow = SimpleNamespace(
        walk_nodes=lambda: iter([("node-1", {"type": "llm"})]),
        user_input_form=lambda: [],
    )
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=workflow))
    monkeypatch.setattr(EvaluationService, "get_supported_metrics", MagicMock(return_value=["answer_relevancy"]))
    validate_embeddings = MagicMock()
    monkeypatch.setattr(EvaluationService, "_validate_default_embedding_model", validate_embeddings)
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[
            DefaultMetric(
                metric="answer_relevancy",
                node_info_list=[NodeInfo(node_id="node-1", type="llm", title="LLM")],
            )
        ],
    )

    EvaluationService.validate_evaluation_config(
        session=session,
        tenant_id="tenant-1",
        target_type="apps",
        target_id="app-1",
        data=config,
    )

    validate_embeddings.assert_called_once_with(tenant_id="tenant-1")


def test_validate_config_does_not_require_embeddings_for_deepeval_metrics(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    workflow = SimpleNamespace(
        walk_nodes=lambda: iter([("node-1", {"type": "llm"})]),
        user_input_form=lambda: [],
    )
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_embeddings.return_value = False
    evaluation_instance.get_supported_metrics.return_value = ["answer_relevancy"]
    monkeypatch.setattr(
        "services.evaluation_service.EvaluationManager.get_evaluation_instance",
        MagicMock(return_value=evaluation_instance),
    )
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=workflow))
    validate_embeddings = MagicMock()
    monkeypatch.setattr(EvaluationService, "_validate_default_embedding_model", validate_embeddings)
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[
            DefaultMetric(
                metric="answer_relevancy",
                node_info_list=[NodeInfo(node_id="node-1", type="llm", title="LLM")],
            )
        ],
    )

    EvaluationService.validate_evaluation_config(
        session=session,
        tenant_id="tenant-1",
        target_type="apps",
        target_id="app-1",
        data=config,
    )

    validate_embeddings.assert_not_called()


def test_validate_knowledge_base_config_uses_retrieval_metrics(monkeypatch: pytest.MonkeyPatch) -> None:
    supported_metrics = MagicMock(return_value=["context_precision"])
    monkeypatch.setattr(EvaluationService, "get_supported_metrics", supported_metrics)
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[DefaultMetric(metric="context_precision", node_info_list=[])],
    )

    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="dataset-1", provider="vendor")

    EvaluationService.validate_evaluation_config(
        session=session,
        tenant_id="tenant-1",
        target_type="knowledge_base",
        target_id="dataset-1",
        data=config,
    )

    supported_metrics.assert_called_once_with(EvaluationCategory.RETRIEVAL)


def test_validate_knowledge_base_config_rejects_workflow_node_scope(monkeypatch: pytest.MonkeyPatch) -> None:
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_embeddings.return_value = False
    monkeypatch.setattr(
        "services.evaluation_service.EvaluationManager.get_evaluation_instance",
        MagicMock(return_value=evaluation_instance),
    )
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[
            DefaultMetric(
                metric="context_precision",
                node_info_list=[
                    NodeInfo(node_id="knowledge-retrieval-1", type="knowledge-retrieval", title="Retriever")
                ],
            )
        ],
    )
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="dataset-1", provider="vendor")

    with pytest.raises(ValueError, match="target-scoped.*must not select workflow nodes"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="knowledge_base",
            target_id="dataset-1",
            data=config,
        )


def test_validate_workflow_retrieval_metric_preserves_node_scope(monkeypatch: pytest.MonkeyPatch) -> None:
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_embeddings.return_value = False
    monkeypatch.setattr(
        "services.evaluation_service.EvaluationManager.get_evaluation_instance",
        MagicMock(return_value=evaluation_instance),
    )
    monkeypatch.setattr(EvaluationService, "get_supported_metrics", MagicMock(return_value=["context_precision"]))
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    workflow = SimpleNamespace(
        walk_nodes=lambda: iter([("knowledge-retrieval-1", {"type": "knowledge-retrieval"})]),
        user_input_form=lambda: [],
    )
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=workflow))
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[
            DefaultMetric(
                metric="context_precision",
                node_info_list=[
                    NodeInfo(node_id="knowledge-retrieval-1", type="knowledge-retrieval", title="Retriever")
                ],
            )
        ],
        judgment_config=JudgmentConfig(
            conditions=[
                JudgmentCondition(
                    variable_selector=["knowledge-retrieval-1", "context_precision"],
                    comparison_operator=">",
                    value="0.5",
                )
            ]
        ),
    )

    EvaluationService.validate_evaluation_config(
        session=session,
        tenant_id="tenant-1",
        target_type="apps",
        target_id="app-1",
        data=config,
    )


def test_validate_knowledge_base_config_rejects_customized_evaluator() -> None:
    config = EvaluationConfigData(
        customized_metrics=CustomizedMetrics(
            evaluation_workflow_id="evaluator-app",
            input_fields={},
            output_fields=[CustomizedMetricOutputField(variable="score", value_type="number")],
        )
    )

    with pytest.raises(ValueError, match="not supported for knowledge bases"):
        EvaluationService.validate_evaluation_config(
            session=MagicMock(),
            tenant_id="tenant-1",
            target_type="knowledge_base",
            target_id="dataset-1",
            data=config,
        )


def test_validate_config_rejects_undeclared_customized_output(monkeypatch: pytest.MonkeyPatch) -> None:
    target = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    evaluator_app = SimpleNamespace(id="evaluator-app")
    session = MagicMock()
    session.scalar.side_effect = [target, evaluator_app]
    target_workflow = SimpleNamespace(walk_nodes=lambda: iter([("node-1", {"type": "llm"})]))
    evaluator_workflow = SimpleNamespace(
        type=WorkflowType.WORKFLOW,
        resolved_kind=WorkflowKind.EVALUATION,
        walk_nodes=lambda: iter([("end", {"type": "end", "outputs": [{"variable": "score", "value_type": "number"}]})]),
    )
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=target_workflow))
    workflow_service = MagicMock()
    workflow_service.get_published_workflow.return_value = evaluator_workflow
    monkeypatch.setattr("services.evaluation_service.WorkflowService", MagicMock(return_value=workflow_service))
    config = EvaluationConfigData(
        customized_metrics=CustomizedMetrics(
            evaluation_workflow_id="evaluator-app",
            input_fields={},
            output_fields=[CustomizedMetricOutputField(variable="missing", value_type="number")],
        )
    )

    with pytest.raises(ValueError, match="not declared.*missing"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="apps",
            target_id="app-1",
            data=config,
        )


def test_validate_config_rejects_non_workflow_app(monkeypatch: pytest.MonkeyPatch) -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="app-1", mode=AppMode.ADVANCED_CHAT.value)
    monkeypatch.setattr(
        EvaluationService,
        "get_supported_metrics",
        MagicMock(return_value=["faithfulness"]),
    )
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[
            DefaultMetric(
                metric="faithfulness",
                node_info_list=[NodeInfo(node_id="node-1", type="llm", title="LLM")],
            )
        ],
    )

    with pytest.raises(ValueError, match="Only workflow apps"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="apps",
            target_id="app-1",
            data=config,
        )


def test_pin_target_workflow_returns_published_snapshot_id(monkeypatch: pytest.MonkeyPatch) -> None:
    target = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    session = MagicMock()
    session.scalar.return_value = target
    workflow = SimpleNamespace(id="workflow-version-7", user_input_form=lambda: [])
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=workflow))

    version_id, resolved_workflow, input_fields = EvaluationService._pin_target_workflow(
        session=session,
        tenant_id="tenant-1",
        target_type="apps",
        target_id="app-1",
    )

    assert version_id == "workflow-version-7"
    assert resolved_workflow is workflow
    assert input_fields == []


def test_extract_llm_context_node_mapping_uses_published_graph_selectors() -> None:
    workflow = SimpleNamespace(
        walk_nodes=lambda: iter(
            [
                ("retrieval-1", {"type": "knowledge-retrieval"}),
                (
                    "llm-1",
                    {
                        "type": "llm",
                        "context": {"enabled": True, "variable_selector": ["retrieval-1", "result"]},
                    },
                ),
            ]
        )
    )

    assert EvaluationService._extract_llm_context_node_mapping(workflow) == {"llm-1": ["retrieval-1"]}


def test_validate_knowledge_base_config_rejects_external_provider() -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="dataset-1", provider="external")
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[DefaultMetric(metric="context_relevance", node_info_list=[])],
    )

    with pytest.raises(ValueError, match="External knowledge bases"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="knowledge_base",
            target_id="dataset-1",
            data=config,
        )


def test_validate_knowledge_base_judgment_uses_dataset_metric_selector(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="dataset-1", provider="vendor")
    monkeypatch.setattr(EvaluationService, "get_supported_metrics", MagicMock(return_value=["context_relevance"]))
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[DefaultMetric(metric="context_relevance", node_info_list=[])],
        judgment_config=JudgmentConfig(
            conditions=[
                JudgmentCondition(
                    variable_selector=["dataset-1", "context_relevance"],
                    comparison_operator=">",
                    value="0.5",
                )
            ]
        ),
    )

    EvaluationService.validate_evaluation_config(
        session=session,
        tenant_id="tenant-1",
        target_type="knowledge_base",
        target_id="dataset-1",
        data=config,
    )


def test_validate_knowledge_base_judgment_rejects_workflow_node_selector(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="dataset-1", provider="vendor")
    monkeypatch.setattr(EvaluationService, "get_supported_metrics", MagicMock(return_value=["context_relevance"]))
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[DefaultMetric(metric="context_relevance", node_info_list=[])],
        judgment_config=JudgmentConfig(
            conditions=[
                JudgmentCondition(
                    variable_selector=["knowledge-retrieval-1", "context_relevance"],
                    comparison_operator=">",
                    value="0.5",
                )
            ]
        ),
    )

    with pytest.raises(ValueError, match="Judgment selector 'knowledge-retrieval-1.context_relevance'"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="knowledge_base",
            target_id="dataset-1",
            data=config,
        )


def test_validate_config_rejects_missing_required_customized_input(monkeypatch: pytest.MonkeyPatch) -> None:
    target = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    evaluator_app = SimpleNamespace(id="evaluator-app")
    session = MagicMock()
    session.scalar.side_effect = [target, evaluator_app]
    target_workflow = SimpleNamespace(walk_nodes=lambda: iter([("node-1", {"type": "llm"})]))
    evaluator_workflow = SimpleNamespace(
        type=WorkflowType.WORKFLOW,
        resolved_kind=WorkflowKind.EVALUATION,
        user_input_form=lambda: [{"variable": "answer", "required": True}],
        walk_nodes=lambda: iter([("end", {"type": "end", "outputs": [{"variable": "score"}]})]),
    )
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=target_workflow))
    workflow_service = MagicMock()
    workflow_service.get_published_workflow.return_value = evaluator_workflow
    monkeypatch.setattr("services.evaluation_service.WorkflowService", MagicMock(return_value=workflow_service))
    config = EvaluationConfigData(
        customized_metrics=CustomizedMetrics(
            evaluation_workflow_id="evaluator-app",
            input_fields={},
            output_fields=[CustomizedMetricOutputField(variable="score", value_type="number")],
        )
    )

    with pytest.raises(ValueError, match="required inputs.*answer"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="apps",
            target_id="app-1",
            data=config,
        )


def test_validate_config_rejects_unknown_customized_selector_node(monkeypatch: pytest.MonkeyPatch) -> None:
    target = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    evaluator_app = SimpleNamespace(id="evaluator-app")
    session = MagicMock()
    session.scalar.side_effect = [target, evaluator_app]
    target_workflow = SimpleNamespace(walk_nodes=lambda: iter([("node-1", {"type": "llm"})]))
    evaluator_workflow = SimpleNamespace(
        type=WorkflowType.WORKFLOW,
        resolved_kind=WorkflowKind.EVALUATION,
        user_input_form=lambda: [{"variable": "answer", "required": True}],
        walk_nodes=lambda: iter([("end", {"type": "end", "outputs": [{"variable": "score"}]})]),
    )
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=target_workflow))
    workflow_service = MagicMock()
    workflow_service.get_published_workflow.return_value = evaluator_workflow
    monkeypatch.setattr("services.evaluation_service.WorkflowService", MagicMock(return_value=workflow_service))
    config = EvaluationConfigData(
        customized_metrics=CustomizedMetrics(
            evaluation_workflow_id="evaluator-app",
            input_fields={"answer": "{{#missing_node.text#}}"},
            output_fields=[CustomizedMetricOutputField(variable="score", value_type="number")],
        )
    )

    with pytest.raises(ValueError, match="missing_node.*published target workflow"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="apps",
            target_id="app-1",
            data=config,
        )


def test_validate_config_rejects_unknown_judgment_selector(monkeypatch: pytest.MonkeyPatch) -> None:
    session = MagicMock()
    session.scalar.return_value = SimpleNamespace(id="app-1", mode=AppMode.WORKFLOW.value)
    workflow = SimpleNamespace(walk_nodes=lambda: iter([("node-1", {"type": "llm"})]))
    monkeypatch.setattr(EvaluationService, "_resolve_published_workflow", MagicMock(return_value=workflow))
    monkeypatch.setattr(EvaluationService, "get_supported_metrics", MagicMock(return_value=["faithfulness"]))
    config = EvaluationConfigData(
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[
            DefaultMetric(
                metric="faithfulness",
                node_info_list=[NodeInfo(node_id="node-1", type="llm", title="LLM")],
            )
        ],
        judgment_config=JudgmentConfig(
            conditions=[
                JudgmentCondition(
                    variable_selector=["node-1", "answer_correctness"],
                    comparison_operator=">",
                    value="0.8",
                )
            ]
        ),
    )

    with pytest.raises(ValueError, match="Judgment selector"):
        EvaluationService.validate_evaluation_config(
            session=session,
            tenant_id="tenant-1",
            target_type="apps",
            target_id="app-1",
            data=config,
        )


def test_cancel_run_stops_each_possible_workflow_task(monkeypatch: pytest.MonkeyPatch) -> None:
    run = SimpleNamespace(
        id="run-1",
        status="running",
        total_items=2,
        celery_task_id="celery-1",
        execution_token="40000000-0000-0000-0000-000000000001",
        heartbeat_at=naive_utc_now(),
        completed_at=None,
        progress_rate=0.4,
    )
    session = MagicMock()
    monkeypatch.setattr(EvaluationService, "get_evaluation_run_detail", MagicMock(return_value=run))
    stop_workflow_tasks = MagicMock()
    monkeypatch.setattr(EvaluationService, "_stop_workflow_tasks", stop_workflow_tasks)
    celery_control = MagicMock()
    monkeypatch.setattr("celery.current_app.control", celery_control)

    EvaluationService.cancel_evaluation_run(
        session=session,
        tenant_id="tenant-1",
        target_type="apps",
        target_id="app-1",
        run_id="run-1",
    )

    stop_workflow_tasks.assert_called_once_with("run-1", 2)
    celery_control.revoke.assert_called_once_with("celery-1", terminate=False)
    assert run.execution_token is None
    assert run.heartbeat_at is None


def test_stop_workflow_tasks_sends_graph_commands_when_queue_stop_flags_fail(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    stop_flag = MagicMock(side_effect=RuntimeError("redis flag failed"))
    graph_manager = MagicMock()
    monkeypatch.setattr(
        "core.app.apps.base_app_queue_manager.AppQueueManager.set_stop_flag_no_user_check",
        stop_flag,
    )
    monkeypatch.setattr("graphon.graph_engine.manager.GraphEngineManager", MagicMock(return_value=graph_manager))

    EvaluationService._stop_workflow_tasks("run-1", 1)

    assert stop_flag.call_count == 2
    assert graph_manager.send_stop_command.call_count == 2


def test_stage_target_cleanup_removes_target_rows_and_preserves_evaluator_defaults(sqlite_session) -> None:
    target_config = EvaluationConfiguration(
        id="10000000-0000-0000-0000-000000000001",
        tenant_id="10000000-0000-0000-0000-000000000010",
        target_type="apps",
        target_id="10000000-0000-0000-0000-000000000100",
        metrics_config=json.dumps({"default_metrics": [], "customized_metrics": None}),
        created_by="10000000-0000-0000-0000-000000001000",
        updated_by="10000000-0000-0000-0000-000000001000",
    )
    referencing_config = EvaluationConfiguration(
        id="20000000-0000-0000-0000-000000000001",
        tenant_id=target_config.tenant_id,
        target_type="apps",
        target_id="20000000-0000-0000-0000-000000000100",
        metrics_config=json.dumps(
            {
                "default_metrics": [{"metric": "faithfulness", "node_info_list": []}],
                "customized_metrics": {"evaluation_workflow_id": target_config.target_id},
            }
        ),
        evaluator_app_id=target_config.target_id,
        created_by=target_config.created_by,
        updated_by=target_config.updated_by,
    )
    run = EvaluationRun(
        id="10000000-0000-0000-0000-000000000002",
        tenant_id=target_config.tenant_id,
        target_type="apps",
        target_id=target_config.target_id,
        target_workflow_version_id="10000000-0000-0000-0000-000000000003",
        evaluation_config_id=target_config.id,
        config_snapshot="{}",
        status=EvaluationRunStatus.COMPLETED,
        total_items=1,
        progress_rate=1.0,
        created_by=target_config.created_by,
    )
    item = EvaluationRunItem(
        id="10000000-0000-0000-0000-000000000004",
        evaluation_run_id=run.id,
        item_index=1,
    )
    sqlite_session.add_all([target_config, referencing_config, run, item])
    sqlite_session.commit()

    EvaluationService.stage_target_evaluation_cleanup(
        session=sqlite_session,
        tenant_id=target_config.tenant_id,
        target_type="apps",
        target_id=target_config.target_id,
    )
    sqlite_session.flush()

    assert sqlite_session.get(EvaluationConfiguration, target_config.id) is None
    assert sqlite_session.get(EvaluationRun, run.id) is None
    assert sqlite_session.get(EvaluationRunItem, item.id) is None
    preserved = sqlite_session.get(EvaluationConfiguration, referencing_config.id)
    assert preserved is not None
    assert preserved.evaluator_app_id is None
    assert preserved.metrics_config_dict["customized_metrics"] is None


@pytest.mark.parametrize("target_type", ["apps", "knowledge_base"])
def test_stage_target_cleanup_deletes_result_file_after_target_deletion_commits(
    sqlite_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    target_type: str,
) -> None:
    tenant_id = "10000000-0000-0000-0000-000000000010"
    target_id = "10000000-0000-0000-0000-000000000100"
    result_file = make_upload_file(
        file_id="10000000-0000-0000-0000-000000000020",
        tenant_id=tenant_id,
        key="evaluation_results/tenant/result.xlsx",
        name="result.xlsx",
        extension="xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        created_by="10000000-0000-0000-0000-000000001000",
        used=True,
    )
    run = EvaluationRun(
        id="10000000-0000-0000-0000-000000000021",
        tenant_id=tenant_id,
        target_type=target_type,
        target_id=target_id,
        evaluation_config_id="10000000-0000-0000-0000-000000000001",
        config_snapshot="{}",
        status=EvaluationRunStatus.COMPLETED,
        result_file_id=result_file.id,
        total_items=1,
        created_by="10000000-0000-0000-0000-000000001000",
    )
    sqlite_session.add_all([result_file, run])
    sqlite_session.commit()
    storage_delete = MagicMock()
    monkeypatch.setattr("extensions.ext_storage.storage.delete", storage_delete)

    EvaluationService.stage_target_evaluation_cleanup(
        session=sqlite_session,
        tenant_id=tenant_id,
        target_type=target_type,
        target_id=target_id,
    )
    sqlite_session.flush()

    assert sqlite_session.get(type(result_file), result_file.id) is None
    storage_delete.assert_not_called()

    sqlite_session.commit()

    storage_delete.assert_called_once_with("evaluation_results/tenant/result.xlsx")


def test_get_run_items_hides_incomplete_target_checkpoints(sqlite_session: Session) -> None:
    run = EvaluationRun(
        id="10000000-0000-0000-0000-000000000012",
        tenant_id="10000000-0000-0000-0000-000000000010",
        target_type="apps",
        target_id="10000000-0000-0000-0000-000000000100",
        target_workflow_version_id="10000000-0000-0000-0000-000000000003",
        evaluation_config_id="10000000-0000-0000-0000-000000000001",
        config_snapshot="{}",
        status=EvaluationRunStatus.RUNNING,
        total_items=2,
        created_by="10000000-0000-0000-0000-000000001000",
    )
    sqlite_session.add_all(
        [
            run,
            EvaluationRunItem(
                evaluation_run_id=run.id,
                item_index=1,
                status=EvaluationRunItemStatus.TARGET_COMPLETED,
            ),
            EvaluationRunItem(
                evaluation_run_id=run.id,
                item_index=2,
                status=EvaluationRunItemStatus.COMPLETED,
            ),
        ]
    )
    sqlite_session.commit()

    items, total = EvaluationService.get_evaluation_run_items(
        sqlite_session,
        run.tenant_id,
        "apps",
        run.target_id,
        run.id,
    )

    assert total == 1
    assert [item.item_index for item in items] == [2]


def test_stage_target_cleanup_does_not_stop_tasks_after_rollback(
    sqlite_session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    config = EvaluationConfiguration(
        id="30000000-0000-0000-0000-000000000001",
        tenant_id="30000000-0000-0000-0000-000000000010",
        target_type="snippets",
        target_id="30000000-0000-0000-0000-000000000100",
        metrics_config=json.dumps({"default_metrics": [], "customized_metrics": None}),
        created_by="30000000-0000-0000-0000-000000001000",
        updated_by="30000000-0000-0000-0000-000000001000",
    )
    run = EvaluationRun(
        id="30000000-0000-0000-0000-000000000002",
        tenant_id=config.tenant_id,
        target_type="snippets",
        target_id=config.target_id,
        target_workflow_version_id="30000000-0000-0000-0000-000000000003",
        evaluation_config_id=config.id,
        config_snapshot="{}",
        status=EvaluationRunStatus.RUNNING,
        total_items=1,
        progress_rate=0.2,
        created_by=config.created_by,
        celery_task_id="celery-task-1",
    )
    sqlite_session.add_all([config, run])
    sqlite_session.commit()
    stop_workflow_tasks = MagicMock()
    monkeypatch.setattr(EvaluationService, "_stop_workflow_tasks", stop_workflow_tasks)

    EvaluationService.stage_target_evaluation_cleanup(
        session=sqlite_session,
        tenant_id=config.tenant_id,
        target_type="snippets",
        target_id=config.target_id,
    )
    sqlite_session.rollback()
    assert sqlite_session.get(EvaluationRun, run.id) is not None

    sqlite_session.commit()

    stop_workflow_tasks.assert_not_called()


def test_retry_failed_run_redispatches_snapshot_and_keeps_checkpoints(
    sqlite_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    tenant = Tenant(name="Evaluation Tenant")
    tenant.id = "30000000-0000-0000-0000-000000000010"
    run_request = EvaluationRunSnapshot(
        file_id="30000000-0000-0000-0000-000000000020",
        evaluation_model_provider="provider",
        evaluation_model="model",
        default_metrics=[DefaultMetric(metric="context_relevance", node_info_list=[])],
        evaluation_framework=EvaluationFrameworkEnum.DEEPEVAL,
        dataset_input_types={"query": "string"},
    )
    run = EvaluationRun(
        id="30000000-0000-0000-0000-000000000021",
        tenant_id=tenant.id,
        target_type="knowledge_base",
        target_id="30000000-0000-0000-0000-000000000030",
        evaluation_config_id="30000000-0000-0000-0000-000000000040",
        config_snapshot=run_request.model_dump_json(),
        status=EvaluationRunStatus.FAILED,
        dataset_file_id=run_request.file_id,
        total_items=2,
        completed_items=0,
        failed_items=2,
        progress_rate=0.9,
        error="worker unavailable",
        celery_task_id="old-task-id",
        created_by="30000000-0000-0000-0000-000000000050",
        started_at=naive_utc_now(),
        completed_at=naive_utc_now(),
    )
    checkpoint = EvaluationRunItem(
        evaluation_run_id=run.id,
        item_index=1,
        status=EvaluationRunItemStatus.COMPLETED,
        inputs='{"query": "first"}',
        judgment="{}",
    )
    sqlite_session.add_all([tenant, run, checkpoint])
    sqlite_session.commit()
    evaluation_instance = MagicMock()
    evaluation_instance.metric_requires_expected_output.return_value = False
    get_evaluation_instance = MagicMock(return_value=evaluation_instance)
    monkeypatch.setattr(
        "services.evaluation_service.EvaluationManager.get_evaluation_instance",
        get_evaluation_instance,
    )
    apply_async = MagicMock()
    monkeypatch.setattr("tasks.evaluation_task.run_evaluation.apply_async", apply_async)

    retried = EvaluationService.retry_evaluation_run(
        session=sqlite_session,
        tenant_id=tenant.id,
        target_type="knowledge_base",
        target_id=run.target_id,
        run_id=run.id,
        dataset_file_content=b"index,query\n1,first\n2,second\n",
        dataset_filename="dataset.csv",
    )

    assert retried.status == EvaluationRunStatus.PENDING
    assert retried.celery_task_id != "old-task-id"
    assert retried.execution_token is None
    assert retried.started_at is None
    assert retried.completed_at is None
    assert retried.error is None
    assert retried.completed_items == 1
    assert retried.failed_items == 0
    assert retried.progress_rate == 0.5
    assert sqlite_session.get(EvaluationRunItem, checkpoint.id) is not None
    apply_async.assert_called_once()
    dispatched_payload = apply_async.call_args.kwargs["args"][0]
    assert dispatched_payload["evaluation_run_id"] == run.id
    assert [item["inputs"] for item in dispatched_payload["input_list"]] == [
        {"query": "first"},
        {"query": "second"},
    ]
    assert dispatched_payload["evaluation_framework"] == "deepeval"
    get_evaluation_instance.assert_called_once_with(EvaluationFrameworkEnum.DEEPEVAL)

    with pytest.raises(ValueError, match="Only failed"):
        EvaluationService.retry_evaluation_run(
            session=sqlite_session,
            tenant_id=tenant.id,
            target_type="knowledge_base",
            target_id=run.target_id,
            run_id=run.id,
            dataset_file_content=b"index,query\n1,first\n2,second\n",
            dataset_filename="dataset.csv",
        )


def test_start_run_persists_and_dispatches_target_workflow_snapshot(monkeypatch: pytest.MonkeyPatch) -> None:
    session = MagicMock()
    session.scalar.return_value = "tenant-1"
    session.query.return_value.filter_by.return_value.filter.return_value.count.return_value = 0

    def assign_run_id(model: object) -> None:
        if isinstance(model, EvaluationRun):
            model.id = "run-1"

    session.add.side_effect = assign_run_id
    monkeypatch.setattr(
        "services.evaluation_service.EvaluationManager.get_configured_framework",
        MagicMock(return_value=EvaluationFrameworkEnum.DEEPEVAL),
    )
    monkeypatch.setattr(
        "services.evaluation_service.EvaluationManager.get_evaluation_instance",
        MagicMock(return_value=MagicMock()),
    )
    monkeypatch.setattr(
        EvaluationService,
        "save_evaluation_config",
        MagicMock(return_value=SimpleNamespace(id="config-1")),
    )
    monkeypatch.setattr(
        EvaluationService,
        "_pin_target_workflow",
        MagicMock(
            return_value=(
                "workflow-version-7",
                SimpleNamespace(walk_nodes=lambda: iter([])),
                [
                    {"variable": "query", "type": "text-input"},
                    {"variable": "enabled", "type": "checkbox"},
                ],
            )
        ),
    )
    apply_async = MagicMock()
    monkeypatch.setattr("tasks.evaluation_task.run_evaluation.apply_async", apply_async)

    run = EvaluationService.start_evaluation_run(
        session=session,
        tenant_id="tenant-1",
        target_type="apps",
        target_id="app-1",
        account_id="account-1",
        dataset_file_content=b"index,query,enabled\n1,hello,true\n",
        dataset_filename="dataset.csv",
        run_request=EvaluationRunRequest(file_id="file-1", evaluation_model="model"),
    )

    assert run.target_workflow_version_id == "workflow-version-7"
    assert run.config_snapshot_dict["evaluation_framework"] == "deepeval"
    assert run.config_snapshot_dict["dataset_input_types"] == {"query": "string", "enabled": "boolean"}
    dispatch_kwargs = apply_async.call_args.kwargs
    assert set(dispatch_kwargs) == {"args", "task_id"}
    assert dispatch_kwargs["args"][0]["evaluation_framework"] == "deepeval"
    dispatched_payload = dispatch_kwargs["args"][0]
    assert dispatched_payload["target_workflow_version_id"] == "workflow-version-7"
    assert dispatched_payload["input_list"][0]["inputs"] == {"query": "hello", "enabled": True}


def test_run_single_app_target_loads_pinned_workflow_version(monkeypatch: pytest.MonkeyPatch) -> None:
    app = SimpleNamespace(id="app-1", tenant_id="tenant-1", mode=AppMode.WORKFLOW.value)
    workflow = SimpleNamespace(id="workflow-version-7")
    session = MagicMock()
    session.scalar.return_value = app
    workflow_service = MagicMock()
    workflow_service.get_published_workflow_by_id.return_value = workflow
    monkeypatch.setattr("services.evaluation_service.WorkflowService", MagicMock(return_value=workflow_service))
    monkeypatch.setattr(
        "core.evaluation.runners.get_service_account_for_app",
        MagicMock(return_value=SimpleNamespace(id="account-1")),
    )
    generator = MagicMock()
    generator.generate.return_value = {"workflow_run_id": "workflow-run-1"}
    monkeypatch.setattr("core.app.apps.workflow.app_generator.WorkflowAppGenerator", MagicMock(return_value=generator))

    result = EvaluationService._run_single_target(
        session=session,
        tenant_id="tenant-1",
        target_type="apps",
        target_id="app-1",
        target_workflow_version_id="workflow-version-7",
        item=EvaluationDatasetInput(index=1, inputs={"query": "hello"}),
        task_id="evaluation:run-1:target:0",
    )

    assert result["workflow_run_id"] == "workflow-run-1"
    workflow_service.get_published_workflow_by_id.assert_called_once_with(
        app_model=app,
        workflow_id="workflow-version-7",
        session=session,
    )
    assert generator.generate.call_args.kwargs["workflow"] is workflow
    assert generator.generate.call_args.kwargs["task_id"] == "evaluation:run-1:target:0"


def test_run_single_app_target_fails_when_pinned_workflow_was_deleted(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    app = SimpleNamespace(id="app-1", tenant_id="tenant-1", mode=AppMode.WORKFLOW.value)
    session = MagicMock()
    session.scalar.return_value = app
    workflow_service = MagicMock()
    workflow_service.get_published_workflow_by_id.return_value = None
    monkeypatch.setattr("services.evaluation_service.WorkflowService", MagicMock(return_value=workflow_service))
    monkeypatch.setattr(
        "core.evaluation.runners.get_service_account_for_app",
        MagicMock(return_value=SimpleNamespace(id="account-1")),
    )
    generator = MagicMock()
    monkeypatch.setattr("core.app.apps.workflow.app_generator.WorkflowAppGenerator", MagicMock(return_value=generator))

    with pytest.raises(EvaluationWorkflowVersionNotFoundError, match="workflow-version-7"):
        EvaluationService._run_single_target(
            session=session,
            tenant_id="tenant-1",
            target_type="apps",
            target_id="app-1",
            target_workflow_version_id="workflow-version-7",
            item=EvaluationDatasetInput(index=1, inputs={"query": "hello"}),
            task_id="evaluation:run-1:target:0",
        )

    generator.generate.assert_not_called()
